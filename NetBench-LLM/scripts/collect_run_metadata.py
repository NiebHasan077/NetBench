#!/usr/bin/env python3
"""
Collect recovered run metadata for a NetBench-LLM model output tree.

This script is intentionally read-only. It scans the current repository for
the evidence needed in a reproducibility package: pipeline settings, training
arguments, LoRA adapter configs, answer/judging/RAG Excel metadata, report
summaries, hashes, git state, and basic environment information.

Example:
    python scripts/collect_run_metadata.py \
        --model-slug qwen3.5-9b \
        --pipeline-script scripts/run_pipeline_specific.sh \
        --output outputs/by_model/qwen3.5-9b/run_metadata.json
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import platform
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]


def rel(path: Path) -> str:
    try:
        return str(path.resolve().relative_to(ROOT.resolve()))
    except ValueError:
        return str(path)


def run_cmd(args: list[str], cwd: Path = ROOT) -> str | None:
    try:
        return subprocess.check_output(
            args,
            cwd=str(cwd),
            stderr=subprocess.DEVNULL,
            text=True,
        ).strip()
    except Exception:
        return None


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            h.update(block)
    return h.hexdigest()


def file_record(path: Path, hash_file: bool = True) -> dict[str, Any]:
    st = path.stat()
    rec: dict[str, Any] = {
        "path": rel(path),
        "bytes": st.st_size,
        "mtime_utc": datetime.fromtimestamp(st.st_mtime, timezone.utc).isoformat(),
    }
    if hash_file and path.is_file():
        rec["sha256"] = sha256_file(path)
    return rec


def parse_shell_assignments(path: Path) -> dict[str, str]:
    """Parse simple VAR=value assignments from the pipeline script."""
    out: dict[str, str] = {}
    if not path.exists():
        return out
    pattern = re.compile(r"^(?:export\s+)?([A-Z][A-Z0-9_]+)=(.*)$")
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        line = line.split("#", 1)[0].strip()
        for segment in line.split(";"):
            segment = segment.strip()
            m = pattern.match(segment)
            if not m:
                continue
            key, raw = m.groups()
            raw = raw.strip()
            if raw.startswith('"') and raw.endswith('"'):
                raw = raw[1:-1]
            elif raw.startswith("'") and raw.endswith("'"):
                raw = raw[1:-1]
            out[key] = raw
    return out


def load_openpyxl():
    try:
        from openpyxl import load_workbook  # type: ignore

        return load_workbook
    except Exception:
        return None


def read_excel_metadata(path: Path) -> dict[str, Any]:
    load_workbook = load_openpyxl()
    if load_workbook is None:
        return {"error": "openpyxl unavailable"}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {"error": str(exc)}
    result: dict[str, Any] = {"sheets": wb.sheetnames}
    if "Metadata" not in wb.sheetnames:
        return result
    meta: dict[str, Any] = {}
    ws = wb["Metadata"]
    for row in ws.iter_rows(values_only=True):
        vals = [v for v in row if v is not None]
        if len(vals) >= 2 and str(vals[0]).lower() != "key":
            meta[str(vals[0])] = vals[1]
    result["metadata"] = meta
    return result


def collect_excel_outputs(model_root: Path) -> dict[str, list[dict[str, Any]]]:
    groups = {
        "answer_files": model_root / "evaluations" / "answers",
        "rag_answer_files": model_root / "evaluations" / "rag_answers",
        "judged_files": model_root / "evaluations" / "judged",
        "report_files": model_root / "evaluations" / "reports",
        "adaptation_reports": model_root / "evaluations" / "adaptation_reports",
    }
    out: dict[str, list[dict[str, Any]]] = {k: [] for k in groups}
    for key, folder in groups.items():
        if not folder.exists():
            continue
        for path in sorted(folder.glob("*")):
            if not path.is_file():
                continue
            rec = file_record(path)
            if path.suffix.lower() == ".xlsx":
                rec["excel"] = read_excel_metadata(path)
            elif path.suffix.lower() == ".json":
                try:
                    rec["json_preview"] = json.loads(path.read_text(encoding="utf-8"))
                except Exception:
                    pass
            out[key].append(rec)
    return out


def collect_parallel_worker_files(model_root: Path) -> list[dict[str, Any]]:
    records = []
    answer_dir = model_root / "evaluations" / "answers"
    if not answer_dir.exists():
        return records
    for path in sorted(answer_dir.glob(".parallel_*/*.json")):
        rec = file_record(path)
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            if isinstance(data, list):
                rec["records"] = len(data)
                if data:
                    rec["first_id"] = data[0].get("id")
                    rec["last_id"] = data[-1].get("id")
        except Exception:
            pass
        records.append(rec)
    return records


def _slug_match(path: Path, model_slug: str) -> bool:
    return model_slug.lower() in str(path).lower()


def collect_training_args(model_slug: str) -> list[dict[str, Any]]:
    try:
        import torch  # type: ignore
    except Exception as exc:
        return [{"error": f"torch unavailable: {exc}"}]

    fields = [
        "output_dir",
        "per_device_train_batch_size",
        "per_device_eval_batch_size",
        "gradient_accumulation_steps",
        "learning_rate",
        "weight_decay",
        "adam_beta1",
        "adam_beta2",
        "adam_epsilon",
        "max_grad_norm",
        "num_train_epochs",
        "max_steps",
        "lr_scheduler_type",
        "warmup_ratio",
        "warmup_steps",
        "logging_steps",
        "eval_steps",
        "save_steps",
        "save_total_limit",
        "bf16",
        "fp16",
        "gradient_checkpointing",
        "optim",
        "deepspeed",
        "dataloader_num_workers",
        "seed",
        "data_seed",
        "report_to",
        "run_name",
    ]
    records = []
    for path in sorted((ROOT / "models").glob("**/training_args.bin")):
        if not _slug_match(path, model_slug):
            continue
        rec: dict[str, Any] = file_record(path)
        try:
            args = torch.load(path, map_location="cpu", weights_only=False)
            values = {}
            for field in fields:
                if hasattr(args, field):
                    value = getattr(args, field)
                    if value is not None:
                        values[field] = str(value) if not isinstance(value, (str, int, float, bool, list, dict)) else value
            rec["training_args"] = values
        except Exception as exc:
            rec["error"] = str(exc)
        records.append(rec)
    return records


def collect_lora_adapters(model_slug: str) -> list[dict[str, Any]]:
    records = []
    fields = [
        "base_model_name_or_path",
        "peft_type",
        "task_type",
        "r",
        "lora_alpha",
        "lora_dropout",
        "bias",
        "target_modules",
        "use_rslora",
        "modules_to_save",
    ]
    for path in sorted((ROOT / "models" / "lora").glob("*/adapter_config.json")):
        if not _slug_match(path, model_slug):
            continue
        rec = file_record(path)
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            rec["adapter_config"] = {field: data.get(field) for field in fields}
        except Exception as exc:
            rec["error"] = str(exc)
        records.append(rec)
    return records


def collect_model_configs(model_slug: str) -> list[dict[str, Any]]:
    fields = [
        "_name_or_path",
        "model_type",
        "architectures",
        "torch_dtype",
        "vocab_size",
        "hidden_size",
        "intermediate_size",
        "num_hidden_layers",
        "num_attention_heads",
        "num_key_value_heads",
        "max_position_embeddings",
        "sliding_window",
        "tie_word_embeddings",
    ]
    records = []
    for path in sorted((ROOT / "models").glob("**/config.json")):
        if not _slug_match(path, model_slug):
            continue
        rec = file_record(path)
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            rec["config"] = {field: data.get(field) for field in fields if field in data}
        except Exception as exc:
            rec["error"] = str(exc)
        records.append(rec)
    return records


def collect_data_files() -> list[dict[str, Any]]:
    candidates = [
        ROOT / "data" / "raw" / "research_corpus_v3.json",
        ROOT / "data" / "raw" / "research_corpus_v3.jsonl",
        ROOT / "data" / "prompts" / "hpn_benchmark_v5.0.jsonl",
        ROOT / "data" / "Instruct-FTD" / "v3_run_plus_json" / "train.jsonl",
        ROOT / "data" / "Instruct-FTD" / "v3_run_plus_json" / "validation.jsonl",
        ROOT / "data" / "Instruct-FTD" / "v3_run_plus_json" / "json_append_report.json",
    ]
    records = []
    for path in candidates:
        if not path.exists():
            continue
        rec = file_record(path)
        if path.suffix == ".jsonl":
            with path.open("r", encoding="utf-8", errors="replace") as f:
                rec["jsonl_lines"] = sum(1 for _ in f)
        elif path.name.endswith(".json"):
            try:
                data = json.loads(path.read_text(encoding="utf-8"))
                if isinstance(data, list):
                    rec["json_records"] = len(data)
                elif isinstance(data, dict):
                    rec["json_keys"] = sorted(data.keys())[:50]
            except Exception:
                pass
        records.append(rec)
    return records


def collect_code_artifacts(rag_dir: Path | None) -> list[dict[str, Any]]:
    paths = [
        ROOT / "training" / "prepare_data.py",
        ROOT / "training" / "prepare_lora_data.py",
        ROOT / "training" / "prepare_instruction_data.py",
        ROOT / "training" / "pretrain_transformers.py",
        ROOT / "training" / "instruction_finetune.py",
        ROOT / "training" / "lora_finetune.py",
        ROOT / "training" / "merge_lora_adapter.py",
        ROOT / "evaluation" / "hpn_qa_benchmark.py",
        ROOT / "evaluation" / "judge_responses.py",
        ROOT / "evaluation" / "benchmark_report.py",
        ROOT / "evaluation" / "lm_adaptation_report.py",
        ROOT / "profiling" / "profile_inference.py",
        ROOT / "profiling" / "profile_rag_inference.py",
        ROOT / "profiling" / "summarize_inference_profiles.py",
        ROOT / "utils" / "model_utils.py",
        ROOT / "requirements.txt",
    ]
    if rag_dir is not None and rag_dir.exists():
        paths.extend(
            [
                rag_dir / "config.yaml",
                rag_dir / "index_corpus.py",
                rag_dir / "query.py",
                rag_dir / "evaluation" / "evaluate_rag.py",
                rag_dir / "evaluation" / "judge_responses.py",
                rag_dir / "evaluation" / "benchmark_report.py",
                rag_dir / "src" / "retriever.py",
                rag_dir / "src" / "generator.py",
                rag_dir / "src" / "store.py",
                rag_dir / "src" / "chunker.py",
                rag_dir / "src" / "embedder.py",
                rag_dir / "src" / "parser.py",
                rag_dir / "requirements.txt",
            ]
        )
    return [file_record(path) for path in paths if path.exists()]


def collect_logs(model_root: Path) -> list[dict[str, Any]]:
    logs_dir = model_root / "logs"
    records = []
    if not logs_dir.exists():
        return records
    interesting = re.compile(
        r"(Model:|Hardware:|Hyperparams:|Full-training GPU workers:|"
        r"Benchmark GPU workers:|Corpus:|RUN |SKIP |Saved:|All done|"
        r"judge|temperature|max_new_tokens|learning_rate|epoch)"
    )
    for path in sorted(logs_dir.glob("*.log")):
        rec = file_record(path, hash_file=True)
        lines = []
        for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
            clean = re.sub(r"\x1b\[[0-9;]*m", "", line)
            if interesting.search(clean):
                lines.append(clean)
        rec["interesting_lines_sample"] = lines[:250]
        rec["interesting_lines_count"] = len(lines)
        records.append(rec)
    return records


def collect_effective_preflight(model_root: Path) -> dict[str, Any]:
    logs_dir = model_root / "logs"
    if not logs_dir.exists():
        return {}
    log_files = sorted(logs_dir.glob("*.log"), key=lambda p: p.stat().st_mtime, reverse=True)
    patterns = {
        "model": r"Model:\s*(.*)$",
        "hardware": r"Hardware:\s*(.*)$",
        "hyperparams": r"Hyperparams:\s*(.*)$",
        "full_training_workers": r"Full-training GPU workers:\s*(.*)$",
        "benchmark_workers": r"Benchmark GPU workers:\s*(.*)$",
        "corpus": r"Corpus:\s*(.*)$",
    }
    for path in log_files:
        values: dict[str, Any] = {"source_log": rel(path)}
        text = path.read_text(encoding="utf-8", errors="replace")
        for raw in text.splitlines():
            line = re.sub(r"\x1b\[[0-9;]*m", "", raw)
            for key, pattern in patterns.items():
                if key in values:
                    continue
                m = re.search(pattern, line)
                if m:
                    values[key] = m.group(1).strip()
        if len(values) > 1:
            return values
    return {}


def collect_rag_config(rag_dir: Path | None) -> dict[str, Any]:
    if rag_dir is None:
        return {}
    config = rag_dir / "config.yaml"
    if not config.exists():
        return {"config_path": str(config), "exists": False}
    rec = file_record(config)
    try:
        import yaml  # type: ignore

        rec["config"] = yaml.safe_load(config.read_text(encoding="utf-8"))
    except Exception as exc:
        rec["error"] = str(exc)
    for extra in ["data/qdrant_store/meta.json", "data/bm25_index.pkl", "data/chunks_cache.pkl", ".index_done"]:
        p = rag_dir / extra
        if p.exists():
            rec.setdefault("artifacts", []).append(file_record(p))
    return rec


def collect_environment() -> dict[str, Any]:
    env = {
        "python": sys.version.replace("\n", " "),
        "platform": platform.platform(),
        "executable": sys.executable,
        "cwd": str(ROOT),
        "collected_at_utc": datetime.now(timezone.utc).isoformat(),
        "git_commit": run_cmd(["git", "rev-parse", "HEAD"]),
        "git_branch": run_cmd(["git", "branch", "--show-current"]),
        "git_status_short": run_cmd(["git", "status", "--short"]),
        "nvidia_smi_query": run_cmd(["nvidia-smi", "--query-gpu=name,memory.total,driver_version", "--format=csv"]),
    }
    try:
        import torch  # type: ignore

        env["torch"] = {
            "version": torch.__version__,
            "cuda": torch.version.cuda,
            "cuda_available": torch.cuda.is_available(),
            "device_count": torch.cuda.device_count() if torch.cuda.is_available() else 0,
        }
    except Exception as exc:
        env["torch_error"] = str(exc)
    return env


def write_table_summary(manifest: dict[str, Any], output_csv: Path) -> None:
    rows = []
    for group_name in ["answer_files", "rag_answer_files", "judged_files"]:
        for rec in manifest.get("outputs", {}).get(group_name, []):
            meta = rec.get("excel", {}).get("metadata", {})
            row = {
                "group": group_name,
                "file": rec.get("path"),
                "model_name": meta.get("model_name"),
                "benchmark_file": meta.get("benchmark_file"),
                "timestamp": meta.get("timestamp"),
                "temperature": meta.get("temperature") or meta.get("generation_temperature"),
                "max_tokens": meta.get("max_new_tokens") or meta.get("generation_max_tokens"),
                "top_p": meta.get("top_p"),
                "judge_model": meta.get("judge_model"),
                "total_questions": meta.get("total_questions"),
                "sha256": rec.get("sha256"),
            }
            rows.append(row)
    if not rows:
        return
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_training_summary(manifest: dict[str, Any], output_csv: Path) -> None:
    rows = []
    for rec in manifest.get("training_args", []):
        args = rec.get("training_args", {})
        rows.append(
            {
                "file": rec.get("path"),
                "output_dir": args.get("output_dir"),
                "train_batch": args.get("per_device_train_batch_size"),
                "eval_batch": args.get("per_device_eval_batch_size"),
                "grad_accum": args.get("gradient_accumulation_steps"),
                "learning_rate": args.get("learning_rate"),
                "epochs": args.get("num_train_epochs"),
                "weight_decay": args.get("weight_decay"),
                "scheduler": args.get("lr_scheduler_type"),
                "warmup_steps": args.get("warmup_steps"),
                "eval_steps": args.get("eval_steps"),
                "save_steps": args.get("save_steps"),
                "optimizer": args.get("optim"),
                "bf16": args.get("bf16"),
                "fp16": args.get("fp16"),
                "gradient_checkpointing": args.get("gradient_checkpointing"),
                "seed": args.get("seed"),
                "sha256": rec.get("sha256"),
            }
        )
    if not rows:
        return
    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Collect recovered NetBench run metadata.")
    parser.add_argument("--model-slug", required=True, help="Model output slug under outputs/by_model/")
    parser.add_argument(
        "--pipeline-script",
        default="scripts/run_pipeline.sh",
        help="Pipeline script used or recovered for this run.",
    )
    parser.add_argument(
        "--rag-dir",
        default="../NetBench-RAG",
        help="Sibling NetBench-RAG directory, if available.",
    )
    parser.add_argument("--output", required=True, help="JSON manifest output path.")
    parser.add_argument(
        "--summary-csv",
        default=None,
        help="Optional CSV summary for answer/judged metadata.",
    )
    parser.add_argument(
        "--training-csv",
        default=None,
        help="Optional CSV summary for saved TrainingArguments metadata.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    model_root = ROOT / "outputs" / "by_model" / args.model_slug
    pipeline_script = ROOT / args.pipeline_script
    rag_dir = (ROOT / args.rag_dir).resolve() if args.rag_dir else None

    manifest: dict[str, Any] = {
        "schema_version": "netbench-run-metadata-v1",
        "model_slug": args.model_slug,
        "environment": collect_environment(),
        "pipeline_script": file_record(pipeline_script) if pipeline_script.exists() else {"path": str(pipeline_script), "exists": False},
        "pipeline_variables": parse_shell_assignments(pipeline_script),
        "code_artifacts": collect_code_artifacts(rag_dir if rag_dir and rag_dir.exists() else None),
        "data_files": collect_data_files(),
        "model_configs": collect_model_configs(args.model_slug),
        "training_args": collect_training_args(args.model_slug),
        "lora_adapters": collect_lora_adapters(args.model_slug),
        "outputs": collect_excel_outputs(model_root),
        "parallel_worker_files": collect_parallel_worker_files(model_root),
        "effective_preflight_from_latest_log": collect_effective_preflight(model_root),
        "logs": collect_logs(model_root),
        "rag": collect_rag_config(rag_dir if rag_dir and rag_dir.exists() else None),
    }

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(manifest, indent=2, sort_keys=True), encoding="utf-8")
    if args.summary_csv:
        write_table_summary(manifest, Path(args.summary_csv))
    if args.training_csv:
        write_training_summary(manifest, Path(args.training_csv))
    print(f"Wrote {output}")
    if args.summary_csv:
        print(f"Wrote {args.summary_csv}")
    if args.training_csv:
        print(f"Wrote {args.training_csv}")


if __name__ == "__main__":
    main()
