#!/usr/bin/env python3
"""Inter-judge agreement report for the HPN-QA benchmark.

Pairs each model's two judged sheets — one per judge (e.g. ``..._by_gpt-5.1.xlsx``
and ``..._by_gemini-3.5-flash.xlsx``) — and quantifies how much the two judges
agree. This supports a *robustness* claim ("model rankings do not depend on the
choice of LLM judge"). It does NOT speak to correctness — agreement between two
LLM judges measures consistency, not ground truth.

Two views:
  • Per-question agreement (pooled over every paired model × question) on the
    four 1-5 dimensions and on ``overall``: Pearson r, Spearman rho, mean
    absolute error, and quadratic-weighted Cohen's kappa (kappa for the integer
    dimensions only; ``overall`` is continuous).
  • System-level (the headline): mean ``overall`` per model under each judge,
    ranked, with Kendall tau and Spearman rho between the two leaderboards.

Outputs:
    outputs/evaluations/reports/HPN_JUDGE_AGREEMENT_{judgeA}_vs_{judgeB}.xlsx
    outputs/evaluations/reports/HPN_JUDGE_AGREEMENT_{judgeA}_vs_{judgeB}.md

Usage:
    python evaluation/judge_agreement.py \\
        --gpt_glob 'outputs/by_model/**/hpn_judged_*_by_gpt-5.1.xlsx' \\
        --gemini_glob 'outputs/by_model/**/hpn_judged_*_by_gemini-3.5-flash.xlsx'

Metrics are hand-rolled in numpy (no scipy dependency); ``_selftest()`` runs on
every invocation and asserts the implementations against known values.
"""

from __future__ import annotations

import argparse
import glob
import os
import re

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# read_judged_excel lives next door and already parses model_name / judge_model
# / the Judged rows out of a Phase-2 sheet. Reuse it rather than reimplementing.
from benchmark_report import read_judged_excel

DIMENSIONS = ["correctness", "completeness", "clarity", "conciseness"]
INT_DIMENSIONS = DIMENSIONS  # the four 1-5 dims are integers; overall is a float
ALL_FIELDS = DIMENSIONS + ["overall"]


# ═══════════════════════════════════════════════════════════════════════
# Hand-rolled statistics (numpy only)
# ═══════════════════════════════════════════════════════════════════════

def _pearson(a, b) -> float:
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    if len(a) < 2 or a.std() == 0 or b.std() == 0:
        return float("nan")
    return float(np.corrcoef(a, b)[0, 1])


def _rankdata(a) -> np.ndarray:
    """Ranks with averaged ties (1-based), matching scipy.stats.rankdata."""
    a = np.asarray(a, dtype=float)
    n = len(a)
    order = np.argsort(a, kind="mergesort")
    ranks = np.empty(n, dtype=float)
    ranks[order] = np.arange(1, n + 1)
    sa = a[order]
    i = 0
    while i < n:
        j = i
        while j + 1 < n and sa[j + 1] == sa[i]:
            j += 1
        if j > i:
            ranks[order[i:j + 1]] = (i + 1 + j + 1) / 2.0
        i = j + 1
    return ranks


def _spearman(a, b) -> float:
    return _pearson(_rankdata(a), _rankdata(b))


def _kendall_tau(a, b) -> float:
    """Kendall tau-b (handles ties); O(n^2), fine for ~63 models."""
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    n = len(a)
    nc = nd = ties_a = ties_b = 0
    for i in range(n):
        for j in range(i + 1, n):
            da = a[i] - a[j]
            db = b[i] - b[j]
            s = da * db
            if s > 0:
                nc += 1
            elif s < 0:
                nd += 1
            else:
                if da == 0:
                    ties_a += 1
                if db == 0:
                    ties_b += 1
    n0 = n * (n - 1) / 2.0
    denom = np.sqrt((n0 - ties_a) * (n0 - ties_b))
    return (nc - nd) / denom if denom > 0 else float("nan")


def _mae(a, b) -> float:
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    return float(np.mean(np.abs(a - b)))


def _quadratic_weighted_kappa(a, b) -> float:
    a = np.asarray(a, dtype=int)
    b = np.asarray(b, dtype=int)
    lo = int(min(a.min(), b.min()))
    hi = int(max(a.max(), b.max()))
    R = hi - lo + 1
    if R < 2:
        return float("nan")
    O = np.zeros((R, R), dtype=float)
    for x, y in zip(a, b):
        O[x - lo, y - lo] += 1
    w = np.zeros((R, R), dtype=float)
    for i in range(R):
        for j in range(R):
            w[i, j] = ((i - j) ** 2) / ((R - 1) ** 2)
    hist_a = O.sum(axis=1)
    hist_b = O.sum(axis=0)
    N = O.sum()
    E = np.outer(hist_a, hist_b) / N
    num = (w * O).sum()
    den = (w * E).sum()
    return float(1 - num / den) if den > 0 else float("nan")


def _selftest() -> None:
    """Assert the hand-rolled metrics against known values on every run."""
    assert abs(_pearson([1, 2, 3, 4], [1, 2, 3, 4]) - 1.0) < 1e-9
    assert abs(_pearson([1, 2, 3, 4], [4, 3, 2, 1]) + 1.0) < 1e-9
    assert abs(_spearman([1, 2, 3, 4], [1, 2, 3, 4]) - 1.0) < 1e-9
    assert abs(_kendall_tau([1, 2, 3, 4], [1, 2, 3, 4]) - 1.0) < 1e-9
    assert abs(_kendall_tau([1, 2, 3, 4], [4, 3, 2, 1]) + 1.0) < 1e-9
    assert abs(_quadratic_weighted_kappa([1, 2, 3, 4], [1, 2, 3, 4]) - 1.0) < 1e-9
    # independent ratings -> kappa ~ 0
    assert abs(_quadratic_weighted_kappa([1, 1, 2, 2], [1, 2, 1, 2])) < 1e-9
    assert abs(_mae([1, 2, 3], [1, 2, 4]) - (1 / 3)) < 1e-9


# ═══════════════════════════════════════════════════════════════════════
# Loading + pairing
# ═══════════════════════════════════════════════════════════════════════

def _model_key(path: str) -> str:
    """Stable model key from a judged filename (judge-suffix independent)."""
    name = os.path.basename(path)
    name = re.sub(r"^hpn_judged_", "", name)
    name = re.sub(r"_by_.*\.xlsx$", "", name)
    return name


def _load_judge_set(pattern: str, excluded: set | None = None,
                    overall_field: str = "overall") -> tuple[dict, str]:
    """Return ({model_key: {id: {field: value}}}, judge_label) for a glob.

    ``overall_field`` selects which sheet column populates the ``overall``
    field: the judge-reported value (default) or the deterministic
    ``overall_formula`` (docs/JUDGE_OVERALL_AUDIT.md)."""
    excluded = excluded or set()
    by_model: dict[str, dict] = {}
    judge_label = "unknown"
    for path in sorted(glob.glob(pattern, recursive=True)):
        _model_name, judge_model, _bench, rows = read_judged_excel(path)
        if judge_model and judge_model != "unknown":
            judge_label = judge_model
        key = _model_key(path)
        per_id: dict = {}
        for r in rows:
            rid = str(r.get("id"))
            if rid in excluded:
                continue
            rec = {}
            ok = True
            for f in ALL_FIELDS:
                v = r.get(overall_field if f == "overall" else f)
                if v is None:
                    ok = False
                    break
                rec[f] = float(v)
            if ok:
                per_id[rid] = rec
        by_model[key] = per_id
    return by_model, judge_label


# ═══════════════════════════════════════════════════════════════════════
# Analysis
# ═══════════════════════════════════════════════════════════════════════

def compute_agreement(set_a: dict, set_b: dict) -> dict:
    common_models = sorted(set(set_a) & set(set_b))

    # --- per-question pooled vectors per field ---
    pooled_a = {f: [] for f in ALL_FIELDS}
    pooled_b = {f: [] for f in ALL_FIELDS}
    n_paired_questions = 0
    for m in common_models:
        ids = set(set_a[m]) & set(set_b[m])
        for rid in ids:
            ra, rb = set_a[m][rid], set_b[m][rid]
            for f in ALL_FIELDS:
                pooled_a[f].append(ra[f])
                pooled_b[f].append(rb[f])
        n_paired_questions += len(ids)

    per_question = {}
    for f in ALL_FIELDS:
        a, b = pooled_a[f], pooled_b[f]
        row = {
            "n": len(a),
            "pearson": _pearson(a, b),
            "spearman": _spearman(a, b),
            "mae": _mae(a, b),
            "qwk": _quadratic_weighted_kappa(a, b) if f in INT_DIMENSIONS else float("nan"),
        }
        per_question[f] = row

    # --- system-level leaderboard on mean overall ---
    leaderboard = []
    for m in common_models:
        mean_a = float(np.mean([set_a[m][i]["overall"] for i in set_a[m]]))
        mean_b = float(np.mean([set_b[m][i]["overall"] for i in set_b[m]]))
        leaderboard.append({"model": m, "mean_a": mean_a, "mean_b": mean_b})

    means_a = [d["mean_a"] for d in leaderboard]
    means_b = [d["mean_b"] for d in leaderboard]
    # ranks: 1 = best (highest mean overall)
    rank_a = {d["model"]: r for r, d in
              enumerate(sorted(leaderboard, key=lambda x: -x["mean_a"]), 1)}
    rank_b = {d["model"]: r for r, d in
              enumerate(sorted(leaderboard, key=lambda x: -x["mean_b"]), 1)}
    for d in leaderboard:
        d["rank_a"] = rank_a[d["model"]]
        d["rank_b"] = rank_b[d["model"]]
        d["rank_delta"] = d["rank_a"] - d["rank_b"]
    leaderboard.sort(key=lambda x: x["rank_a"])

    system = {
        "kendall_tau": _kendall_tau(means_a, means_b),
        "spearman": _spearman(means_a, means_b),
        "pearson": _pearson(means_a, means_b),
        "n_models": len(common_models),
    }

    return {
        "common_models": common_models,
        "n_models": len(common_models),
        "n_paired_questions": n_paired_questions,
        "per_question": per_question,
        "leaderboard": leaderboard,
        "system": system,
    }


# ═══════════════════════════════════════════════════════════════════════
# Reporting
# ═══════════════════════════════════════════════════════════════════════

def _fmt(x) -> str:
    return "n/a" if x is None or (isinstance(x, float) and np.isnan(x)) else f"{x:.3f}"


def write_markdown(result: dict, label_a: str, label_b: str, path: str) -> None:
    s = result["system"]
    lines = []
    lines.append(f"# HPN-QA Inter-Judge Agreement: `{label_a}` vs `{label_b}`\n")
    lines.append(
        f"Paired **{result['n_models']} models** "
        f"({result['n_paired_questions']} question-judgements). Both judges scored "
        f"identical answers at `temperature=0.0`.\n")
    lines.append("> Agreement measures *consistency* between the two judges, not "
                 "correctness. It supports a robustness claim only.\n")

    lines.append("## System-level (leaderboard robustness)\n")
    lines.append("| Metric | Value |")
    lines.append("|---|---|")
    lines.append(f"| Kendall tau-b (model ranking) | **{_fmt(s['kendall_tau'])}** |")
    lines.append(f"| Spearman rho (model ranking)  | **{_fmt(s['spearman'])}** |")
    lines.append(f"| Pearson r (mean overall)      | {_fmt(s['pearson'])} |")
    lines.append(f"| Models compared               | {s['n_models']} |\n")

    lines.append("## Per-question agreement (pooled)\n")
    lines.append("| Dimension | n | Pearson | Spearman | MAE | Quadratic-weighted kappa |")
    lines.append("|---|---|---|---|---|---|")
    for f in ALL_FIELDS:
        r = result["per_question"][f]
        lines.append(f"| {f} | {r['n']} | {_fmt(r['pearson'])} | "
                     f"{_fmt(r['spearman'])} | {_fmt(r['mae'])} | {_fmt(r['qwk'])} |")
    lines.append("\n_kappa is reported for the four integer 1-5 dimensions; "
                 "`overall` is a continuous weighted average._\n")

    lines.append("## Per-model leaderboard (mean overall under each judge)\n")
    lines.append(f"| Model | {label_a} mean | rank | {label_b} mean | rank | rank Δ |")
    lines.append("|---|---|---|---|---|---|")
    for d in result["leaderboard"]:
        lines.append(f"| {d['model']} | {d['mean_a']:.3f} | {d['rank_a']} | "
                     f"{d['mean_b']:.3f} | {d['rank_b']} | {d['rank_delta']:+d} |")
    lines.append("")

    with open(path, "w") as fh:
        fh.write("\n".join(lines))


def write_excel(result: dict, label_a: str, label_b: str, path: str) -> None:
    wb = Workbook()
    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")

    def _header(ws, headers):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c, value=h)
            cell.font = bold
            cell.fill = head_fill

    ws = wb.active
    ws.title = "Summary"
    _header(ws, ["Metric", "Value"])
    s = result["system"]
    summary_rows = [
        ("judge_a", label_a),
        ("judge_b", label_b),
        ("models_compared", s["n_models"]),
        ("paired_question_judgements", result["n_paired_questions"]),
        ("kendall_tau_b_ranking", round(s["kendall_tau"], 4)),
        ("spearman_rho_ranking", round(s["spearman"], 4)),
        ("pearson_r_mean_overall", round(s["pearson"], 4)),
    ]
    for i, (k, v) in enumerate(summary_rows, 2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 24

    ws2 = wb.create_sheet("PerQuestionAgreement")
    _header(ws2, ["dimension", "n", "pearson", "spearman", "mae", "qwk"])
    for i, f in enumerate(ALL_FIELDS, 2):
        r = result["per_question"][f]
        ws2.cell(row=i, column=1, value=f)
        ws2.cell(row=i, column=2, value=r["n"])
        ws2.cell(row=i, column=3, value=round(r["pearson"], 4))
        ws2.cell(row=i, column=4, value=round(r["spearman"], 4))
        ws2.cell(row=i, column=5, value=round(r["mae"], 4))
        ws2.cell(row=i, column=6, value=(None if np.isnan(r["qwk"]) else round(r["qwk"], 4)))
    ws2.column_dimensions["A"].width = 16

    ws3 = wb.create_sheet("Leaderboard")
    _header(ws3, ["model", f"{label_a}_mean", f"{label_a}_rank",
                  f"{label_b}_mean", f"{label_b}_rank", "rank_delta"])
    for i, d in enumerate(result["leaderboard"], 2):
        ws3.cell(row=i, column=1, value=d["model"])
        ws3.cell(row=i, column=2, value=round(d["mean_a"], 4))
        ws3.cell(row=i, column=3, value=d["rank_a"])
        ws3.cell(row=i, column=4, value=round(d["mean_b"], 4))
        ws3.cell(row=i, column=5, value=d["rank_b"])
        ws3.cell(row=i, column=6, value=d["rank_delta"])
    ws3.column_dimensions["A"].width = 52

    wb.save(path)


# ═══════════════════════════════════════════════════════════════════════
# CLI
# ═══════════════════════════════════════════════════════════════════════

def _sanitize(label: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "-", label)


def main() -> None:
    _selftest()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--gpt_glob",
        default="outputs/by_model/**/hpn_judged_*_by_gpt-5.1.xlsx",
        help="Glob for judge A (primary) judged sheets.")
    parser.add_argument(
        "--gemini_glob",
        default="outputs/by_model/**/hpn_judged_*_by_gemini-3.5-flash.xlsx",
        help="Glob for judge B (second opinion) judged sheets.")
    parser.add_argument(
        "--output_dir", default="outputs/evaluations/reports",
        help="Directory for the agreement report (xlsx + md).")
    parser.add_argument(
        "--excluded", default="",
        help="Optional CSV with an 'id' column of audit-removed question IDs to "
             "exclude, matching the scoring set (analysis/excluded_items.csv).")
    parser.add_argument(
        "--overall_field", default="overall",
        choices=["overall", "overall_formula"],
        help="Which composite drives the 'overall' rows and the system-level "
             "leaderboard: the judge-reported value (default) or the "
             "deterministic rubric aggregate (docs/JUDGE_OVERALL_AUDIT.md). "
             "With overall_formula the output stem gains a _formula suffix.")
    args = parser.parse_args()

    excluded: set[str] = set()
    if args.excluded and os.path.exists(args.excluded):
        import csv
        with open(args.excluded) as fh:
            excluded = {row["id"].strip() for row in csv.DictReader(fh)}
        print(f"   excluding {len(excluded)} audited items from agreement")

    set_a, label_a = _load_judge_set(args.gpt_glob, excluded, args.overall_field)
    set_b, label_b = _load_judge_set(args.gemini_glob, excluded, args.overall_field)
    if not set_a or not set_b:
        raise SystemExit(
            f"❌ No judged sheets matched.\n   A ({label_a}): {len(set_a)} files\n"
            f"   B ({label_b}): {len(set_b)} files\n   check the globs.")

    result = compute_agreement(set_a, set_b)

    os.makedirs(args.output_dir, exist_ok=True)
    stem = f"HPN_JUDGE_AGREEMENT_{_sanitize(label_a)}_vs_{_sanitize(label_b)}"
    if args.overall_field != "overall":
        stem += "_formula"
    md_path = os.path.join(args.output_dir, stem + ".md")
    xlsx_path = os.path.join(args.output_dir, stem + ".xlsx")
    write_markdown(result, label_a, label_b, md_path)
    write_excel(result, label_a, label_b, xlsx_path)

    s = result["system"]
    print(f"✅ Paired {result['n_models']} models, "
          f"{result['n_paired_questions']} question-judgements.")
    print(f"   Kendall tau-b (ranking) : {_fmt(s['kendall_tau'])}")
    print(f"   Spearman rho (ranking)  : {_fmt(s['spearman'])}")
    print(f"   overall Pearson r       : {_fmt(result['per_question']['overall']['pearson'])}")
    print(f"   → {md_path}")
    print(f"   → {xlsx_path}")


if __name__ == "__main__":
    main()
