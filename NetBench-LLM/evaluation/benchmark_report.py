#!/usr/bin/env python3
"""
HPN Q&A Benchmark — Comparative Report (Phase 3)
==================================================

Reads two or more Phase 2 judged Excel files and produces a side-by-side
comparative benchmark report in both Excel and Markdown.

Outputs:
    outputs/evaluations/reports/hpn_benchmark_report_by_{judge}.xlsx   — Multi-sheet workbook
    outputs/evaluations/reports/HPN_BENCHMARK_REPORT_by_{judge}.md     — Markdown report

Excel sheets:
    • **Comparison**      — one row per question, side-by-side overall scores
    • **Overall Summary** — avg / min / max per dimension per model
    • **By Category**     — avg scores per category per model (pivot)
    • **By Difficulty**   — avg scores per difficulty per model
    • **Head-to-Head**    — win / tie / loss counts (overall & by category)

Usage:
    # Compare two judged files
    python evaluation/benchmark_report.py \\
        --judged_files outputs/evaluations/hpn_judged_Llama-3.1-8B-base-instruct.xlsx \\
                       outputs/evaluations/hpn_judged_Llama-3.1-8B-trained-new-instruct.xlsx

    # Custom output directory
    python evaluation/benchmark_report.py \\
        --judged_files outputs/evaluations/hpn_judged_*.xlsx \\
        --output_dir outputs/evaluations
"""

import argparse
import os
import re
import sys
import warnings
from datetime import datetime
from pathlib import Path

# Add project root to path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# ═══════════════════════════════════════════════════════════════════════
# Constants
# ═══════════════════════════════════════════════════════════════════════

DEFAULT_OUTPUT_DIR = "outputs/evaluations/reports"
DIMS = ["correctness", "completeness", "clarity", "conciseness", "overall"]
DIFFICULTIES = ["easy", "medium", "hard"]


def _benchmark_tag(benchmark_path: str) -> str:
    """Extract a concise tag from the benchmark filename.

    Examples:
        'hpn_qa_benchmark_v4_general_skills.json' → 'v4_general_skills'
        'hpn_benchmark_v5.0_all.jsonl'            → 'v5.0_all'
        'hpn_qa_benchmark.json'                   → ''  (no tag)
    """
    stem = Path(benchmark_path).stem
    for prefix in ("hpn_qa_benchmark", "hpn_benchmark"):
        if stem.startswith(prefix) and len(stem) > len(prefix):
            return stem[len(prefix) + 1:]
    return ""

# ═══════════════════════════════════════════════════════════════════════
# Excel style constants (consistent with Phase 1 & 2)
# ═══════════════════════════════════════════════════════════════════════

_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_META_KEY_FONT = Font(name="Calibri", bold=True, size=11)
_SECTION_FONT = Font(name="Calibri", bold=True, size=12)
_TITLE_FONT = Font(name="Calibri", bold=True, size=14)

_WINNER_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_LOSER_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_TIE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def _style_header_row_at(ws, row_num: int, num_cols: int):
    """Apply header styling at an arbitrary row number."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER


def _apply_border_row(ws, row_num: int, num_cols: int):
    """Apply thin border to a data row."""
    for col in range(1, num_cols + 1):
        ws.cell(row=row_num, column=col).border = _THIN_BORDER


# ═══════════════════════════════════════════════════════════════════════
# Read Phase 2 judged Excel file
# ═══════════════════════════════════════════════════════════════════════

def read_judged_excel(filepath: str) -> tuple[str, str, str, list[dict]]:
    """Read a Phase 2 judged Excel file.

    Returns:
        (model_name, judge_model, benchmark_file, list_of_judged_row_dicts)
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)

    # Get model name, judge model, and benchmark from Metadata sheet
    model_name = Path(filepath).stem
    judge_model = "unknown"
    benchmark_file = ""
    if "Metadata" in wb.sheetnames:
        ws_meta = wb["Metadata"]
        for row in ws_meta.iter_rows(min_row=2, max_col=2, values_only=True):
            if row[0] is not None:
                key = str(row[0])
                val = str(row[1]) if row[1] else ""
                if key == "model_name":
                    model_name = val or model_name
                elif key == "judge_model":
                    judge_model = val or judge_model
                elif key == "benchmark_file":
                    benchmark_file = val

    # Read Judged sheet
    ws = wb["Judged"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = {}
        for col_idx, header in enumerate(headers):
            record[header] = row[col_idx] if col_idx < len(row) else None
        # Only include rows with valid overall score
        if record.get("overall") is not None:
            rows.append(record)

    wb.close()
    return model_name, judge_model, benchmark_file, rows


# ═══════════════════════════════════════════════════════════════════════
# Aggregation helpers
# ═══════════════════════════════════════════════════════════════════════

def _avg(rows: list[dict], key: str) -> float:
    vals = [float(r[key]) for r in rows if r.get(key) is not None]
    return round(sum(vals) / len(vals), 2) if vals else 0.0


def _min_val(rows: list[dict], key: str) -> float:
    vals = [float(r[key]) for r in rows if r.get(key) is not None]
    return round(min(vals), 2) if vals else 0.0


def _max_val(rows: list[dict], key: str) -> float:
    vals = [float(r[key]) for r in rows if r.get(key) is not None]
    return round(max(vals), 2) if vals else 0.0


# ═══════════════════════════════════════════════════════════════════════
# Excel report builder
# ═══════════════════════════════════════════════════════════════════════

def build_excel_report(
    model_data: dict[str, list[dict]],
    output_path: str,
):
    """Build the multi-sheet comparative Excel report.

    Args:
        model_data: dict mapping model_name → list of judged row dicts.
        output_path: path for the output .xlsx file.
    """
    model_names = list(model_data.keys())
    wb = Workbook()

    # ── Sheet 1: Comparison ─────────────────────────────────────────
    _build_comparison_sheet(wb, model_data, model_names)

    # ── Sheet 2: Overall Summary ────────────────────────────────────
    _build_overall_summary_sheet(wb, model_data, model_names)

    # ── Sheet 3: By Category ───────────────────────────────────────
    _build_by_category_sheet(wb, model_data, model_names)

    # ── Sheet 4: By Difficulty ──────────────────────────────────────
    _build_by_difficulty_sheet(wb, model_data, model_names)

    # ── Sheet 5: Head-to-Head (only if exactly 2 models) ───────────
    if len(model_names) == 2:
        _build_head_to_head_sheet(wb, model_data, model_names)

    # Remove the default empty sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"💾 Excel report saved: {output_path}")


def _build_comparison_sheet(wb, model_data, model_names):
    """Sheet with one row per question, side-by-side overall scores."""
    ws = wb.create_sheet("Comparison", 0)

    # Build a lookup: question_id → {model_name: row_dict}
    all_ids = []
    id_meta = {}       # id → {category, difficulty, question}
    id_scores = {}     # id → {model_name: {dim: score}}

    for mname, rows in model_data.items():
        for r in rows:
            qid = str(r.get("id", ""))
            if qid not in id_scores:
                id_scores[qid] = {}
                all_ids.append(qid)
                id_meta[qid] = {
                    "category": r.get("category", ""),
                    "difficulty": r.get("difficulty", ""),
                    "question": r.get("question", ""),
                }
            id_scores[qid][mname] = {d: r.get(d) for d in DIMS}

    # Headers: id | category | difficulty | question | model_A_overall | model_B_overall | delta | winner
    headers = ["id", "category", "difficulty", "question"]
    for mname in model_names:
        headers.append(f"{mname}_overall")
    if len(model_names) == 2:
        headers.extend(["delta", "winner"])

    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row_at(ws, 1, len(headers))

    for row_idx, qid in enumerate(all_ids, start=2):
        meta = id_meta[qid]
        ws.cell(row=row_idx, column=1, value=qid)
        ws.cell(row=row_idx, column=2, value=meta["category"])
        ws.cell(row=row_idx, column=3, value=meta["difficulty"])
        ws.cell(row=row_idx, column=4, value=meta["question"])

        scores_for_q = []
        for m_idx, mname in enumerate(model_names):
            overall = id_scores[qid].get(mname, {}).get("overall")
            col = 5 + m_idx
            ws.cell(row=row_idx, column=col, value=overall)
            ws.cell(row=row_idx, column=col).alignment = _CENTER_ALIGN
            scores_for_q.append(overall)

        if len(model_names) == 2:
            s_a, s_b = scores_for_q
            delta_col = 5 + len(model_names)
            winner_col = delta_col + 1

            if s_a is not None and s_b is not None:
                delta = round(float(s_b) - float(s_a), 2)
                ws.cell(row=row_idx, column=delta_col, value=delta)
                ws.cell(row=row_idx, column=delta_col).alignment = _CENTER_ALIGN

                if delta > 0:
                    winner = model_names[1]
                    ws.cell(row=row_idx, column=winner_col).fill = _WINNER_FILL
                elif delta < 0:
                    winner = model_names[0]
                    ws.cell(row=row_idx, column=winner_col).fill = _LOSER_FILL
                else:
                    winner = "TIE"
                    ws.cell(row=row_idx, column=winner_col).fill = _TIE_FILL

                ws.cell(row=row_idx, column=winner_col, value=winner)
                ws.cell(row=row_idx, column=winner_col).alignment = _CENTER_ALIGN

        _apply_border_row(ws, row_idx, len(headers))

    # Column widths
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 55
    for i in range(len(model_names)):
        ws.column_dimensions[get_column_letter(5 + i)].width = 30
    if len(model_names) == 2:
        ws.column_dimensions[get_column_letter(5 + len(model_names))].width = 10
        ws.column_dimensions[get_column_letter(6 + len(model_names))].width = 35

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(all_ids) + 1}"


def _build_overall_summary_sheet(wb, model_data, model_names):
    """Avg / Min / Max per dimension per model."""
    ws = wb.create_sheet("Overall Summary")

    # Headers: Dimension | model_A_avg | model_A_min | model_A_max | model_B_avg | ...
    headers = ["Dimension"]
    for mname in model_names:
        headers.extend([f"{mname}_avg", f"{mname}_min", f"{mname}_max"])

    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row_at(ws, 1, len(headers))

    for row_off, dim in enumerate(DIMS, start=2):
        ws.cell(row=row_off, column=1, value=dim)
        ws.cell(row=row_off, column=1).font = _META_KEY_FONT
        for m_idx, mname in enumerate(model_names):
            rows = model_data[mname]
            base_col = 2 + m_idx * 3
            ws.cell(row=row_off, column=base_col, value=_avg(rows, dim))
            ws.cell(row=row_off, column=base_col + 1, value=_min_val(rows, dim))
            ws.cell(row=row_off, column=base_col + 2, value=_max_val(rows, dim))
            for c in range(base_col, base_col + 3):
                ws.cell(row=row_off, column=c).alignment = _CENTER_ALIGN

        _apply_border_row(ws, row_off, len(headers))

    # Highlight best avg for each dim (green fill)
    if len(model_names) >= 2:
        for row_off, dim in enumerate(DIMS, start=2):
            avgs = []
            for m_idx, mname in enumerate(model_names):
                avgs.append((_avg(model_data[mname], dim), m_idx))
            best_idx = max(avgs, key=lambda x: x[0])[1]
            best_col = 2 + best_idx * 3
            ws.cell(row=row_off, column=best_col).fill = _WINNER_FILL

    ws.column_dimensions["A"].width = 18
    for i in range(1, len(headers)):
        ws.column_dimensions[get_column_letter(i + 1)].width = 20


def _build_by_category_sheet(wb, model_data, model_names):
    """Avg scores per category per model."""
    ws = wb.create_sheet("By Category")

    # Collect all categories
    all_categories = set()
    for rows in model_data.values():
        for r in rows:
            all_categories.add(r.get("category", "Unknown"))
    categories = sorted(all_categories)

    # Headers: Category | model_A_correctness | model_A_completeness | ... | model_A_overall | model_B_... | ...
    headers = ["Category"]
    for mname in model_names:
        for dim in DIMS:
            headers.append(f"{mname}_{dim}")

    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row_at(ws, 1, len(headers))

    for row_off, cat in enumerate(categories, start=2):
        ws.cell(row=row_off, column=1, value=cat)
        ws.cell(row=row_off, column=1).font = _META_KEY_FONT

        for m_idx, mname in enumerate(model_names):
            cat_rows = [r for r in model_data[mname]
                        if r.get("category") == cat]
            for d_idx, dim in enumerate(DIMS):
                col = 2 + m_idx * len(DIMS) + d_idx
                ws.cell(row=row_off, column=col, value=_avg(cat_rows, dim))
                ws.cell(row=row_off, column=col).alignment = _CENTER_ALIGN

        _apply_border_row(ws, row_off, len(headers))

    # Highlight best overall per category
    if len(model_names) >= 2:
        overall_offsets = [m_idx * len(DIMS) + DIMS.index("overall")
                          for m_idx in range(len(model_names))]
        for row_off, cat in enumerate(categories, start=2):
            best_val = -1
            best_col = -1
            for m_idx, offset in enumerate(overall_offsets):
                col = 2 + offset
                val = ws.cell(row=row_off, column=col).value
                if val is not None and val > best_val:
                    best_val = val
                    best_col = col
            if best_col > 0:
                ws.cell(row=row_off, column=best_col).fill = _WINNER_FILL

    ws.column_dimensions["A"].width = 35
    for i in range(1, len(headers)):
        ws.column_dimensions[get_column_letter(i + 1)].width = 16


def _build_by_difficulty_sheet(wb, model_data, model_names):
    """Avg scores per difficulty per model."""
    ws = wb.create_sheet("By Difficulty")

    headers = ["Difficulty"]
    for mname in model_names:
        for dim in DIMS:
            headers.append(f"{mname}_{dim}")

    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row_at(ws, 1, len(headers))

    for row_off, diff in enumerate(DIFFICULTIES, start=2):
        ws.cell(row=row_off, column=1, value=diff)
        ws.cell(row=row_off, column=1).font = _META_KEY_FONT

        for m_idx, mname in enumerate(model_names):
            diff_rows = [r for r in model_data[mname]
                         if str(r.get("difficulty", "")).lower() == diff]
            for d_idx, dim in enumerate(DIMS):
                col = 2 + m_idx * len(DIMS) + d_idx
                ws.cell(row=row_off, column=col, value=_avg(diff_rows, dim))
                ws.cell(row=row_off, column=col).alignment = _CENTER_ALIGN

        _apply_border_row(ws, row_off, len(headers))

    # Highlight best overall per difficulty
    if len(model_names) >= 2:
        overall_offsets = [m_idx * len(DIMS) + DIMS.index("overall")
                          for m_idx in range(len(model_names))]
        for row_off in range(2, 2 + len(DIFFICULTIES)):
            best_val = -1
            best_col = -1
            for m_idx, offset in enumerate(overall_offsets):
                col = 2 + offset
                val = ws.cell(row=row_off, column=col).value
                if val is not None and val > best_val:
                    best_val = val
                    best_col = col
            if best_col > 0:
                ws.cell(row=row_off, column=best_col).fill = _WINNER_FILL

    ws.column_dimensions["A"].width = 14
    for i in range(1, len(headers)):
        ws.column_dimensions[get_column_letter(i + 1)].width = 16


def _build_head_to_head_sheet(wb, model_data, model_names):
    """Win / Tie / Loss counts — only for exactly 2 models."""
    ws = wb.create_sheet("Head-to-Head")
    m_a, m_b = model_names

    rows_a = {str(r.get("id", "")): r for r in model_data[m_a]}
    rows_b = {str(r.get("id", "")): r for r in model_data[m_b]}
    common_ids = sorted(set(rows_a.keys()) & set(rows_b.keys()))

    # ── Section 1: Overall head-to-head ─────────────────────────────
    ws.cell(row=1, column=1, value="Overall Head-to-Head")
    ws.cell(row=1, column=1).font = _TITLE_FONT

    h2h_headers = ["Metric", f"{m_a} Wins", "Ties", f"{m_b} Wins",
                    "Total", f"{m_a} Win%", f"{m_b} Win%"]
    for col_idx, h in enumerate(h2h_headers, start=1):
        ws.cell(row=2, column=col_idx, value=h)
    _style_header_row_at(ws, 2, len(h2h_headers))

    for row_off, dim in enumerate(DIMS, start=3):
        a_wins = 0
        b_wins = 0
        ties = 0
        for qid in common_ids:
            sa = rows_a[qid].get(dim)
            sb = rows_b[qid].get(dim)
            if sa is not None and sb is not None:
                if float(sa) > float(sb):
                    a_wins += 1
                elif float(sb) > float(sa):
                    b_wins += 1
                else:
                    ties += 1

        total = a_wins + b_wins + ties
        ws.cell(row=row_off, column=1, value=dim)
        ws.cell(row=row_off, column=1).font = _META_KEY_FONT
        ws.cell(row=row_off, column=2, value=a_wins)
        ws.cell(row=row_off, column=3, value=ties)
        ws.cell(row=row_off, column=4, value=b_wins)
        ws.cell(row=row_off, column=5, value=total)
        ws.cell(row=row_off, column=6, value=f"{a_wins / max(total, 1) * 100:.1f}%")
        ws.cell(row=row_off, column=7, value=f"{b_wins / max(total, 1) * 100:.1f}%")

        for c in range(2, 8):
            ws.cell(row=row_off, column=c).alignment = _CENTER_ALIGN

        # Highlight the winner column
        if a_wins > b_wins:
            ws.cell(row=row_off, column=2).fill = _WINNER_FILL
        elif b_wins > a_wins:
            ws.cell(row=row_off, column=4).fill = _WINNER_FILL
        else:
            ws.cell(row=row_off, column=3).fill = _TIE_FILL

        _apply_border_row(ws, row_off, len(h2h_headers))

    # ── Section 2: Head-to-head by category ─────────────────────────
    cat_start = 3 + len(DIMS) + 1
    ws.cell(row=cat_start, column=1, value="Head-to-Head by Category (Overall Dimension)")
    ws.cell(row=cat_start, column=1).font = _TITLE_FONT

    all_categories = sorted(set(
        r.get("category", "Unknown")
        for rows in model_data.values() for r in rows
    ))

    cat_headers = ["Category", f"{m_a} Wins", "Ties", f"{m_b} Wins",
                   f"{m_a} Avg", f"{m_b} Avg", "Delta"]
    cat_start += 1
    for col_idx, h in enumerate(cat_headers, start=1):
        ws.cell(row=cat_start, column=col_idx, value=h)
    _style_header_row_at(ws, cat_start, len(cat_headers))

    for cat_off, cat in enumerate(all_categories, start=1):
        r_idx = cat_start + cat_off
        a_wins = 0
        b_wins = 0
        ties = 0
        a_vals = []
        b_vals = []

        for qid in common_ids:
            if rows_a[qid].get("category") != cat:
                continue
            sa = rows_a[qid].get("overall")
            sb = rows_b[qid].get("overall")
            if sa is not None and sb is not None:
                a_vals.append(float(sa))
                b_vals.append(float(sb))
                if float(sa) > float(sb):
                    a_wins += 1
                elif float(sb) > float(sa):
                    b_wins += 1
                else:
                    ties += 1

        a_avg = round(sum(a_vals) / len(a_vals), 2) if a_vals else 0.0
        b_avg = round(sum(b_vals) / len(b_vals), 2) if b_vals else 0.0
        delta = round(b_avg - a_avg, 2)

        ws.cell(row=r_idx, column=1, value=cat)
        ws.cell(row=r_idx, column=1).font = _META_KEY_FONT
        ws.cell(row=r_idx, column=2, value=a_wins)
        ws.cell(row=r_idx, column=3, value=ties)
        ws.cell(row=r_idx, column=4, value=b_wins)
        ws.cell(row=r_idx, column=5, value=a_avg)
        ws.cell(row=r_idx, column=6, value=b_avg)
        ws.cell(row=r_idx, column=7, value=delta)

        for c in range(2, 8):
            ws.cell(row=r_idx, column=c).alignment = _CENTER_ALIGN

        if a_wins > b_wins:
            ws.cell(row=r_idx, column=2).fill = _WINNER_FILL
        elif b_wins > a_wins:
            ws.cell(row=r_idx, column=4).fill = _WINNER_FILL

        _apply_border_row(ws, r_idx, len(cat_headers))

    # Column widths
    ws.column_dimensions["A"].width = 35
    for i in range(1, max(len(h2h_headers), len(cat_headers))):
        ws.column_dimensions[get_column_letter(i + 1)].width = 18


# ═══════════════════════════════════════════════════════════════════════
# Markdown report builder
# ═══════════════════════════════════════════════════════════════════════

def build_markdown_report(
    model_data: dict[str, list[dict]],
    output_path: str,
    judge_model: str = "unknown",
):
    """Build a comparative Markdown report.

    Args:
        model_data: dict mapping model_name → list of judged row dicts.
        output_path: path for the output .md file.
        judge_model: name of the LLM used as judge.
    """
    model_names = list(model_data.keys())
    lines = []

    def ln(s=""):
        lines.append(s)

    # ── Title ───────────────────────────────────────────────────────
    ln("# HPN Q&A Benchmark — Comparative Report")
    ln()
    ln(f"**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    ln(f"**Judge model:** `{judge_model}`")
    ln(f"**Models compared:** {len(model_names)}")
    for i, m in enumerate(model_names, 1):
        n_scored = len(model_data[m])
        ln(f"  {i}. `{m}` ({n_scored} scored questions)")
    ln()

    # ── Executive Summary ───────────────────────────────────────────
    ln("## Executive Summary")
    ln()

    # Overall averages table
    ln("| Dimension | " + " | ".join(f"`{m}`" for m in model_names) + " |")
    ln("| --- | " + " | ".join("---:" for _ in model_names) + " |")
    for dim in DIMS:
        row_vals = []
        for mname in model_names:
            row_vals.append(f"{_avg(model_data[mname], dim):.2f}")
        ln(f"| **{dim.capitalize()}** | " + " | ".join(row_vals) + " |")
    ln()

    # Winner declaration (if 2 models)
    if len(model_names) == 2:
        avg_a = _avg(model_data[model_names[0]], "overall")
        avg_b = _avg(model_data[model_names[1]], "overall")
        delta = round(avg_b - avg_a, 2)
        if delta > 0:
            winner = model_names[1]
            diff_sign = f"+{delta}"
        elif delta < 0:
            winner = model_names[0]
            diff_sign = f"{delta}"
        else:
            winner = "TIE"
            diff_sign = "0.00"

        ln(f"> **Overall winner: `{winner}`** (delta = {diff_sign} on the 1–5 scale)")
        ln()

    # ── Scores by Category ──────────────────────────────────────────
    ln("## Scores by Category")
    ln()

    all_categories = sorted(set(
        r.get("category", "Unknown")
        for rows in model_data.values() for r in rows
    ))

    ln("| Category | " + " | ".join(f"`{m}` overall" for m in model_names) + " |")
    ln("| --- | " + " | ".join("---:" for _ in model_names) + " |")

    for cat in all_categories:
        row_vals = []
        for mname in model_names:
            cat_rows = [r for r in model_data[mname]
                        if r.get("category") == cat]
            row_vals.append(f"{_avg(cat_rows, 'overall'):.2f}")
        ln(f"| {cat} | " + " | ".join(row_vals) + " |")
    ln()

    # ── Scores by Difficulty ────────────────────────────────────────
    ln("## Scores by Difficulty")
    ln()

    ln("| Difficulty | " + " | ".join(f"`{m}` overall" for m in model_names) + " |")
    ln("| --- | " + " | ".join("---:" for _ in model_names) + " |")

    for diff in DIFFICULTIES:
        row_vals = []
        for mname in model_names:
            diff_rows = [r for r in model_data[mname]
                         if str(r.get("difficulty", "")).lower() == diff]
            row_vals.append(f"{_avg(diff_rows, 'overall'):.2f}")
        ln(f"| {diff} | " + " | ".join(row_vals) + " |")
    ln()

    # ── Head-to-Head (if 2 models) ─────────────────────────────────
    if len(model_names) == 2:
        m_a, m_b = model_names
        rows_a = {str(r.get("id", "")): r for r in model_data[m_a]}
        rows_b = {str(r.get("id", "")): r for r in model_data[m_b]}
        common_ids = sorted(set(rows_a.keys()) & set(rows_b.keys()))

        ln("## Head-to-Head")
        ln()
        ln(f"Comparing per-question scores across {len(common_ids)} common questions.")
        ln()
        ln(f"| Dimension | `{m_a}` Wins | Ties | `{m_b}` Wins |")
        ln("| --- | ---: | ---: | ---: |")

        for dim in DIMS:
            a_w = b_w = t = 0
            for qid in common_ids:
                sa = rows_a[qid].get(dim)
                sb = rows_b[qid].get(dim)
                if sa is not None and sb is not None:
                    if float(sa) > float(sb):
                        a_w += 1
                    elif float(sb) > float(sa):
                        b_w += 1
                    else:
                        t += 1
            ln(f"| **{dim.capitalize()}** | {a_w} | {t} | {b_w} |")
        ln()

        # Category-level H2H
        ln("### Head-to-Head by Category (Overall)")
        ln()
        ln(f"| Category | `{m_a}` Wins | Ties | `{m_b}` Wins | `{m_a}` Avg | `{m_b}` Avg | Δ |")
        ln("| --- | ---: | ---: | ---: | ---: | ---: | ---: |")

        for cat in all_categories:
            a_w = b_w = t = 0
            a_vals = []
            b_vals = []
            for qid in common_ids:
                if rows_a[qid].get("category") != cat:
                    continue
                sa = rows_a[qid].get("overall")
                sb = rows_b[qid].get("overall")
                if sa is not None and sb is not None:
                    a_vals.append(float(sa))
                    b_vals.append(float(sb))
                    if float(sa) > float(sb):
                        a_w += 1
                    elif float(sb) > float(sa):
                        b_w += 1
                    else:
                        t += 1
            a_avg = round(sum(a_vals) / len(a_vals), 2) if a_vals else 0.0
            b_avg = round(sum(b_vals) / len(b_vals), 2) if b_vals else 0.0
            d = round(b_avg - a_avg, 2)
            d_str = f"+{d}" if d > 0 else f"{d}"
            ln(f"| {cat} | {a_w} | {t} | {b_w} | {a_avg:.2f} | {b_avg:.2f} | {d_str} |")
        ln()

    # ── Per-Dimension Analysis ──────────────────────────────────────
    ln("## Per-Dimension Analysis")
    ln()

    for dim in DIMS:
        if dim == "overall":
            continue
        ln(f"### {dim.capitalize()}")
        ln()
        ln(f"| Model | Avg | Min | Max |")
        ln("| --- | ---: | ---: | ---: |")
        for mname in model_names:
            rows = model_data[mname]
            ln(f"| `{mname}` | {_avg(rows, dim):.2f} | "
               f"{_min_val(rows, dim):.2f} | {_max_val(rows, dim):.2f} |")
        ln()

    # ── Scoring Rubric Reminder ─────────────────────────────────────
    ln("## Scoring Rubric")
    ln()
    ln("Each answer was scored by `" + judge_model + "` on four dimensions (1–5 scale):")
    ln()
    ln("| Dimension | Weight | Description |")
    ln("| --- | ---: | --- |")
    ln("| Correctness | 40% | Technical accuracy relative to reference |")
    ln("| Completeness | 30% | Coverage of key concepts |")
    ln("| Clarity | 20% | Organization and readability |")
    ln("| Conciseness | 10% | Focus without unnecessary padding |")
    ln()
    ln("**Overall** = Correctness×0.4 + Completeness×0.3 + Clarity×0.2 + Conciseness×0.1")
    ln()

    # Write
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"📝 Markdown report saved: {output_path}")


# ═══════════════════════════════════════════════════════════════════════
# Main pipeline
# ═══════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="HPN Q&A Benchmark — Comparative Report (Phase 3)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  # Compare two judged files\n"
            "  python evaluation/benchmark_report.py \\\n"
            "      --judged_files outputs/evaluations/judged/hpn_judged_Llama-3.1-8B-base-instruct_by_gemini-2.5-flash.xlsx \\\n"
            "                     outputs/evaluations/judged/hpn_judged_Llama-3.1-8B-trained-new-instruct_by_gemini-2.5-flash.xlsx\n"
            "\n"
            "  # Custom output directory\n"
            "  python evaluation/benchmark_report.py \\\n"
            "      --judged_files outputs/evaluations/judged/hpn_judged_*.xlsx \\\n"
            "      --output_dir outputs/evaluations/reports\n"
        ),
    )

    parser.add_argument(
        "--judged_files",
        type=str,
        nargs="+",
        required=True,
        help="Path(s) to Phase 2 judged Excel file(s) "
             "(e.g. outputs/evaluations/hpn_judged_*.xlsx)",
    )
    parser.add_argument(
        "--output_dir",
        type=str,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for output files (default: {DEFAULT_OUTPUT_DIR})",
    )

    args = parser.parse_args()

    # ── Banner ──────────────────────────────────────────────────────
    print("\n" + "=" * 70)
    print("  📊  HPN Q&A Benchmark — Comparative Report (Phase 3)")
    print("=" * 70)

    # ── Load judged files ───────────────────────────────────────────
    model_data = {}
    judge_models_seen = set()
    benchmark_files_seen = set()
    for jf in args.judged_files:
        path = Path(jf)
        if not path.exists():
            print(f"❌ File not found: {jf}")
            continue
        model_name, judge_model, benchmark_file, rows = read_judged_excel(jf)
        model_data[model_name] = rows
        judge_models_seen.add(judge_model)
        if benchmark_file:
            benchmark_files_seen.add(benchmark_file)
        print(f"   📋 Loaded {len(rows)} scored rows from: {model_name} "
              f"(judged by {judge_model})")

    if len(model_data) < 2:
        print("⚠️  Need at least 2 judged files for a comparative report.")
        if len(model_data) == 1:
            print("   (Single-model summaries are already in the Phase 2 output.)")
        sys.exit(1)

    # Derive a judge tag for filenames
    if len(judge_models_seen) == 1:
        judge_model_name = judge_models_seen.pop()
    else:
        judge_model_name = "_".join(sorted(judge_models_seen))
    judge_tag = re.sub(r'[^\w\-.]', '_', judge_model_name)

    # Derive benchmark tag for filenames
    bmk_tag = ""
    if len(benchmark_files_seen) == 1:
        bmk_tag = _benchmark_tag(benchmark_files_seen.pop())
    bmk_suffix = f"_{bmk_tag}" if bmk_tag else ""

    print(f"\n⚙️  Configuration:")
    print(f"   Models        : {len(model_data)}")
    for i, (m, rows) in enumerate(model_data.items(), 1):
        print(f"     {i}. {m} ({len(rows)} questions)")
    print(f"   Judge model   : {judge_model_name}")
    print(f"   Output dir    : {args.output_dir}")

    # ── Build reports ───────────────────────────────────────────────
    excel_path = os.path.join(
        args.output_dir, f"hpn_benchmark_report{bmk_suffix}_by_{judge_tag}.xlsx"
    )
    md_path = os.path.join(
        args.output_dir, f"HPN_BENCHMARK_REPORT{bmk_suffix}_by_{judge_tag}.md"
    )

    print()
    build_excel_report(model_data, excel_path)
    build_markdown_report(model_data, md_path, judge_model=judge_model_name)

    # ── Final summary ───────────────────────────────────────────────
    print(f"\n{'=' * 70}")
    print("🏁 Comparative report generated!")
    print(f"   📊 Excel : {excel_path}")
    print(f"   📝 Markdown: {md_path}")
    print(f"{'=' * 70}\n")


if __name__ == "__main__":
    main()
