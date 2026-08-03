#!/usr/bin/env python3
"""Merge per-GPU benchmark answer shard workbooks.

The direct LLM and RAG answer generators both write Excel files with
``Metadata`` and ``Answers`` sheets.  This helper combines shard workbooks into
one normal Phase-1 answer file so the existing judge/report scripts can consume
parallel benchmark output without changes.
"""

from __future__ import annotations

import argparse
import json
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_META_KEY_FONT = Font(name="Calibri", bold=True, size=11)


def _read_workbook(path: Path) -> tuple[dict[str, Any], list[str], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True)

    metadata: dict[str, Any] = {}
    if "Metadata" in wb.sheetnames:
        ws_meta = wb["Metadata"]
        for row in ws_meta.iter_rows(min_row=2, max_col=2, values_only=True):
            if row[0] is not None:
                metadata[str(row[0])] = row[1]

    if "Answers" not in wb.sheetnames:
        wb.close()
        raise ValueError(f"{path} has no Answers sheet")

    ws_ans = wb["Answers"]
    headers = [cell.value for cell in next(ws_ans.iter_rows(min_row=1, max_row=1))]
    headers = [str(h) for h in headers if h is not None]

    rows: list[dict[str, Any]] = []
    for row in ws_ans.iter_rows(min_row=2, values_only=True):
        record = {}
        for idx, header in enumerate(headers):
            record[header] = row[idx] if idx < len(row) else None
        if record.get("id") is not None:
            rows.append(record)

    wb.close()
    return metadata, headers, rows


def _question_order(benchmark_path: str | None) -> dict[str, int]:
    if not benchmark_path:
        return {}
    path = Path(benchmark_path)
    if not path.exists():
        return {}

    if path.suffix.lower() == ".jsonl":
        questions = []
        with open(path, encoding="utf-8") as f:
            for line in f:
                if line.strip():
                    questions.append(json.loads(line))
    else:
        with open(path, encoding="utf-8") as f:
            try:
                data = json.load(f)
            except json.JSONDecodeError:
                f.seek(0)
                questions = [json.loads(line) for line in f if line.strip()]
            else:
                questions = data.get("questions", data if isinstance(data, list) else [])

    return {str(q.get("id", "")): idx for idx, q in enumerate(questions)}


def _style_header_row(ws, num_cols: int) -> None:
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER


def _auto_width(ws, headers: list[str]) -> None:
    wide = {
        "question",
        "reference_answer",
        "model_answer",
        "retrieved_chunks",
        "reranker_scores",
        "keywords",
    }
    for idx, header in enumerate(headers, start=1):
        letter = get_column_letter(idx)
        if header in wide:
            ws.column_dimensions[letter].width = 55
        else:
            ws.column_dimensions[letter].width = min(max(len(header) + 4, 12), 28)


def merge_shards(
    shard_files: list[Path],
    output_file: Path,
    benchmark_path: str | None = None,
) -> None:
    if not shard_files:
        raise ValueError("No shard files provided")

    merged_by_id: dict[str, dict[str, Any]] = {}
    metadata: dict[str, Any] = {}
    headers: list[str] = []
    total_generation_time = 0.0

    for shard in shard_files:
        shard_meta, shard_headers, rows = _read_workbook(shard)
        if not metadata:
            metadata = dict(shard_meta)
        for header in shard_headers:
            if header not in headers:
                headers.append(header)
        for row in rows:
            merged_by_id[str(row["id"])] = row
            try:
                total_generation_time += float(row.get("generation_time_sec") or 0.0)
            except (TypeError, ValueError):
                pass

    order = _question_order(benchmark_path)
    rows = list(merged_by_id.values())
    rows.sort(key=lambda r: (order.get(str(r.get("id", "")), 10**9), str(r.get("id", ""))))

    metadata.update(
        {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "total_questions": len(rows),
            "total_generation_time_sec": round(total_generation_time, 3),
            "parallel_shards": len(shard_files),
            "source_shards": ", ".join(str(p) for p in shard_files),
        }
    )
    if benchmark_path:
        metadata["benchmark_file"] = benchmark_path

    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws_meta = wb.active
    ws_meta.title = "Metadata"
    ws_meta.append(["Key", "Value"])
    for key, value in metadata.items():
        ws_meta.append([key, value])
    _style_header_row(ws_meta, 2)
    ws_meta.column_dimensions["A"].width = 28
    ws_meta.column_dimensions["B"].width = 90
    for row in ws_meta.iter_rows(min_row=2):
        row[0].font = _META_KEY_FONT
        row[1].alignment = _WRAP_ALIGN

    ws_ans = wb.create_sheet("Answers")
    ws_ans.append(headers)
    _style_header_row(ws_ans, len(headers))
    for row_data in rows:
        ws_ans.append([row_data.get(header) for header in headers])
    for row in ws_ans.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = _WRAP_ALIGN
            cell.border = _THIN_BORDER
    ws_ans.freeze_panes = "A2"
    ws_ans.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    _auto_width(ws_ans, headers)

    wb.save(output_file)
    print(f"Merged {len(shard_files)} shard(s), {len(rows)} answer(s) -> {output_file}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Merge benchmark answer shard Excel files.")
    parser.add_argument("--shard_files", nargs="+", required=True)
    parser.add_argument("--output_file", required=True)
    parser.add_argument("--benchmark", default=None)
    args = parser.parse_args()

    merge_shards(
        shard_files=[Path(p) for p in args.shard_files],
        output_file=Path(args.output_file),
        benchmark_path=args.benchmark,
    )


if __name__ == "__main__":
    main()
