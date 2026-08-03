#!/usr/bin/env python3
"""Remove the empty `source_paper` column from released answer/judged workbooks.

The column was emitted by every answer/judge writer but never populated (the
2026-07-18 scan found 0 non-empty cells across all 173 released workbooks), so
it is dropped from the released sheets; the writers no longer emit it. Each
workbook records the operation in its Metadata sheet.

Safety: a workbook is only modified if every `source_paper` cell in its data
sheet is empty; otherwise it is reported and left untouched. Idempotent:
workbooks without the column are skipped. Cell values and styles shift left
with the deletion; column widths and the auto-filter range are re-mapped.

Usage (from the repo root):
    python NetBench-LLM/evaluation/drop_empty_source_paper.py \
        --by-model NetBench-LLM/outputs/by_model [--dry-run]
"""
from __future__ import annotations
import argparse
import glob
import os
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

COL = "source_paper"
NOTE_KEY = "source_paper_removed_note"
NOTE = ("empty source_paper column removed by "
        "evaluation/drop_empty_source_paper.py ({}); the column was never "
        "populated in any released answers/judged sheet")


def drop(path: str, dry_run: bool = False) -> str:
    wb = load_workbook(path)
    sheet = next((s for s in ("Judged", "Answers") if s in wb.sheetnames), None)
    if sheet is None:
        wb.close()
        return "ERROR no data sheet"
    ws = wb[sheet]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if COL not in headers:
        wb.close()
        return "skip (already absent)"
    col_idx = headers.index(COL) + 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=col_idx).value
        if v not in (None, ""):
            wb.close()
            return f"ERROR non-empty {COL} at row {r} ({v!r})"
    if dry_run:
        wb.close()
        return "would drop"

    widths = {column_index_from_string(k): v.width
              for k, v in ws.column_dimensions.items() if v.width}
    ws.delete_cols(col_idx)
    for k in list(ws.column_dimensions.keys()):
        del ws.column_dimensions[k]
    for i, w in sorted(widths.items()):
        if i == col_idx:
            continue
        ws.column_dimensions[get_column_letter(i - (1 if i > col_idx else 0))].width = w
    if ws.auto_filter.ref:
        first, last = ws.auto_filter.ref.split(":")
        row_end = "".join(ch for ch in last if ch.isdigit())
        old_end = column_index_from_string("".join(ch for ch in last if ch.isalpha()))
        ws.auto_filter.ref = f"{first}:{get_column_letter(old_end - 1)}{row_end}"

    meta = wb["Metadata"]
    mrow = meta.max_row + 1
    for col, value in ((1, NOTE_KEY), (2, NOTE.format(date.today().isoformat()))):
        c = meta.cell(row=mrow, column=col, value=value)
        prev = meta.cell(row=mrow - 1, column=col)
        c.border = copy(prev.border)
        if col == 1:
            c.font = copy(prev.font)

    wb.save(path)
    wb.close()
    return f"dropped (was col {get_column_letter(col_idx)}, {sheet})"


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--by-model", default="outputs/by_model",
                    help="by_model tree with */evaluations/{answers,judged}/*.xlsx")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    paths = sorted(glob.glob(os.path.join(args.by_model, "*/evaluations/answers/*.xlsx"))
                   + glob.glob(os.path.join(args.by_model, "*/evaluations/judged/*.xlsx")))
    if not paths:
        raise SystemExit(f"no workbooks under {args.by_model}")
    counts: dict[str, int] = {}
    for p in paths:
        res = drop(p, dry_run=args.dry_run)
        key = res.split(" (")[0]
        counts[key] = counts.get(key, 0) + 1
        if key.startswith("ERROR"):
            print(f"{res:34s} {p}")
    print("summary:", ", ".join(f"{k}: {v}" for k, v in sorted(counts.items())),
          f"({len(paths)} workbooks)")


if __name__ == "__main__":
    main()
