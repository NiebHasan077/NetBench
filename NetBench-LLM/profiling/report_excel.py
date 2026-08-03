#!/usr/bin/env python3
"""
Excel report module for profiling results.

Generates .xlsx workbooks with auto-sized columns, formatted headers,
and number formatting from profiling JSON data.

Output:
  - per-run training workbook  (Config, Dataset, Loss Curve, GPU Monitor, Memory)
  - per-run inference workbook  (Config, Prompts, Output Scaling, Input Scaling, Summary)
  - cumulative workbook         (Training Summary, Inference Summary, comparison sheets)
"""

import os
import logging
from typing import List, Dict, Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, numbers
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# ── Style constants ─────────────────────────────────────────────────
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2196F3", end_color="2196F3",
                           fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
ALT_ROW_FILL = PatternFill(start_color="F5F5F5", end_color="F5F5F5",
                            fill_type="solid")
NUM_FMT_2 = "0.00"
NUM_FMT_4 = "0.0000"
NUM_FMT_INT = "#,##0"
NUM_FMT_SCI = "0.00E+00"


# ── Helpers ─────────────────────────────────────────────────────────

def _auto_width(ws, min_width: int = 10, max_width: int = 50):
    """Auto-fit column widths based on content."""
    for col_cells in ws.columns:
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(
            max(max_len + 2, min_width), max_width
        )


def _write_header(ws, headers: List[str], row: int = 1):
    """Write formatted header row."""
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def _write_row(ws, row: int, values: list, fmt_map: Optional[Dict[int, str]] = None):
    """Write a data row with optional number formats. col indices are 1-based."""
    for col, val in enumerate(values, 1):
        # Convert non-scalar types to strings so openpyxl can handle them
        if isinstance(val, (list, tuple, dict)):
            val = str(val)
        cell = ws.cell(row=row, column=col, value=val)
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if fmt_map and col in fmt_map and val is not None:
            cell.number_format = fmt_map[col]
        if row % 2 == 0:
            cell.fill = ALT_ROW_FILL


def _kv_sheet(ws, title: str, kv_pairs: List[tuple]):
    """Write key-value pairs in two columns."""
    ws.title = title
    _write_header(ws, ["Parameter", "Value"])
    for i, (k, v) in enumerate(kv_pairs, 2):
        # Convert non-scalar types to strings
        if isinstance(v, (list, tuple, dict)):
            v = str(v)
        _write_row(ws, i, [k, v])
    _auto_width(ws)


def _save_wb(wb: Workbook, path: str):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    logger.info(f"    📄  Saved: {path}")


# ═══════════════════════════════════════════════════════════════════
# 1.  PER-RUN TRAINING WORKBOOK
# ═══════════════════════════════════════════════════════════════════

def write_training_run_excel(report: Dict, out_path: str):
    """One .xlsx per training run: Config, Dataset, Loss, GPU, Memory sheets."""
    wb = Workbook()

    # ── Config sheet ───────────────────────────────────────────────
    ws = wb.active
    cfg = report.get("config", {})
    sys_ = report.get("system", {})
    timing = report.get("timing", {})
    kv = [
        ("Run Name", report.get("run_name", "")),
        ("Mode", report.get("mode", "")),
        ("Timestamp", report.get("timestamp", "")),
        ("Model Path", cfg.get("model_path", "")),
        ("Model Name", cfg.get("model_name", "")),
        ("Dtype", cfg.get("dtype", "")),
        ("Num Epochs", cfg.get("num_epochs", "")),
        ("Max Steps", cfg.get("max_steps", "")),
        ("Batch Size", cfg.get("batch_size", "")),
        ("Gradient Accumulation", cfg.get("grad_accum", "")),
        ("Effective Batch Size", cfg.get("effective_batch", "")),
        ("Learning Rate", cfg.get("lr", "")),
        ("Warmup Ratio", cfg.get("warmup_ratio", "")),
        ("Weight Decay", cfg.get("weight_decay", "")),
        ("Max Grad Norm", cfg.get("max_grad_norm", "")),
        ("Optimizer", cfg.get("optimizer", "")),
        ("LR Scheduler", cfg.get("lr_scheduler", "")),
        ("Gradient Checkpointing", cfg.get("grad_ckpt", "")),
        ("", ""),
        ("Hostname", sys_.get("hostname", "")),
        ("OS", sys_.get("os", "")),
        ("Python", sys_.get("python", "")),
        ("PyTorch", sys_.get("torch", "")),
        ("CUDA", sys_.get("cuda", "")),
        ("", ""),
        ("Model Load (s)", timing.get("model_load_s", "")),
        ("Total Training (s)", timing.get("total_train_s", "")),
        ("Total Training (min)", timing.get("total_train_min", "")),
        ("Per-Step (s)", timing.get("per_step_s", "")),
        ("Per-Epoch (s)", timing.get("per_epoch_s", "")),
        ("Total Steps", timing.get("total_steps", "")),
        ("Model Save (s)", timing.get("model_save_s", "")),
    ]
    _kv_sheet(ws, "Config", kv)

    # ── Dataset sheet ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Dataset")
    ds = report.get("dataset", {})
    seq_stats = ds.get("sequence_length_stats", {})
    kv2 = [
        ("Data Dir", ds.get("data_dir", "")),
        ("Dataset Type", ds.get("dataset_type", "")),
        ("Source File", ds.get("source_file", "")),
        ("Description", ds.get("description", "")),
        ("Format", ds.get("format", "")),
        ("Tokenizer", ds.get("tokenizer", "")),
        ("Train Samples", ds.get("train_samples", "")),
        ("Eval Samples", ds.get("eval_samples", "")),
        ("Train Tokens", ds.get("train_tokens", "")),
    ]
    for k, v in seq_stats.items():
        kv2.append((f"Seq Length — {k}", v))
    _kv_sheet(ws2, "Dataset", kv2)

    # ── Loss Curve sheet ───────────────────────────────────────────
    ws3 = wb.create_sheet("Loss Curve")
    curve = report.get("training_loss", {}).get("loss_curve", [])
    if curve:
        headers = ["Step", "Epoch", "Loss", "LR", "Wall (s)",
                    "GPU Alloc (MB)", "GPU Resv (MB)"]
        _write_header(ws3, headers)
        fmt = {3: NUM_FMT_4, 4: NUM_FMT_SCI, 5: NUM_FMT_2, 6: NUM_FMT_INT, 7: NUM_FMT_INT}
        for i, pt in enumerate(curve, 2):
            _write_row(ws3, i, [
                pt.get("step"), pt.get("epoch"), pt.get("loss"),
                pt.get("lr"), pt.get("wall_s"),
                pt.get("gpu_alloc_mb"), pt.get("gpu_resv_mb"),
            ], fmt)
    _auto_width(ws3)

    # ── GPU Monitor sheet ──────────────────────────────────────────
    ws4 = wb.create_sheet("GPU Monitor")
    ts = report.get("gpu_monitor_timeseries", [])
    if ts:
        headers = ["Time (s)", "GPU", "Utilization (%)", "Mem Used (MB)",
                    "Mem Total (MB)", "Temp (°C)", "Power (W)"]
        _write_header(ws4, headers)
        t0 = ts[0]["t"]
        fmt = {1: NUM_FMT_2, 3: NUM_FMT_2, 4: NUM_FMT_INT, 5: NUM_FMT_INT,
               6: NUM_FMT_2, 7: NUM_FMT_2}
        for i, s in enumerate(ts, 2):
            _write_row(ws4, i, [
                s["t"] - t0, s["gpu"], s["util_pct"],
                s["mem_used_mb"], s["mem_total_mb"],
                s["temp_c"], s["power_w"],
            ], fmt)
    _auto_width(ws4)

    # ── Memory sheet ───────────────────────────────────────────────
    ws5 = wb.create_sheet("Memory")
    mem = report.get("memory", {})
    kv5 = [
        ("Peak GPU VRAM (MB)", mem.get("peak_gpu_vram_mb", "")),
        ("CPU RAM Before (MB)", mem.get("cpu_ram_before_mb", "")),
        ("CPU RAM After (MB)", mem.get("cpu_ram_after_mb", "")),
        ("CPU RAM Delta (MB)", mem.get("cpu_ram_delta_mb", "")),
    ]
    for gid, peak in mem.get("per_gpu_peak_mb", {}).items():
        kv5.append((f"GPU {gid} Peak (MB)", peak))
    _kv_sheet(ws5, "Memory", kv5)

    _save_wb(wb, out_path)


# ═══════════════════════════════════════════════════════════════════
# 2.  PER-RUN INFERENCE WORKBOOK
# ═══════════════════════════════════════════════════════════════════

def write_inference_run_excel(report: Dict, out_path: str):
    """One .xlsx per inference run: Config, Prompts, Scaling, Summary."""
    wb = Workbook()

    # ── Config sheet ───────────────────────────────────────────────
    ws = wb.active
    sys_ = report.get("system", {})
    ml = report.get("model_load", {})
    kv = [
        ("Model Path", report.get("model_path", "")),
        ("Model Name", report.get("model_name", "")),
        ("Device", report.get("device", "")),
        ("Timestamp", report.get("timestamp", "")),
        ("", ""),
        ("Hostname", sys_.get("hostname", "")),
        ("OS", sys_.get("os", "")),
        ("Python", sys_.get("python", "")),
        ("PyTorch", sys_.get("torch", "")),
        ("CUDA", sys_.get("cuda", "")),
        ("", ""),
        ("Model Load Time (s)", ml.get("load_time_s", "")),
        ("VRAM After Load (MB)", ml.get("gpu_vram_after_load_mb", "")),
    ]
    _kv_sheet(ws, "Config", kv)

    # ── Prompt Benchmarks sheet ────────────────────────────────────
    ws2 = wb.create_sheet("Prompt Benchmarks")
    prompts = report.get("benchmark_prompts", [])
    if prompts:
        headers = ["Name", "Preview", "Max New Tokens", "Num Runs",
                    "Avg Time (s)", "Avg Tok/s", "Avg ms/tok",
                    "Avg TTFT (s)", "Peak VRAM (MB)"]
        _write_header(ws2, headers)
        fmt = {5: NUM_FMT_2, 6: NUM_FMT_2, 7: NUM_FMT_2, 8: NUM_FMT_4, 9: NUM_FMT_INT}
        for i, p in enumerate(prompts, 2):
            _write_row(ws2, i, [
                p.get("name"), p.get("instruction_preview", "")[:80],
                p.get("max_new_tokens"), p.get("num_runs"),
                p.get("avg_total_time_s"), p.get("avg_tokens_per_second"),
                p.get("avg_ms_per_token"), p.get("avg_ttft_s"),
                p.get("peak_gpu_vram_mb"),
            ], fmt)
    _auto_width(ws2, max_width=60)

    # ── Output Scaling sheet ───────────────────────────────────────
    ws3 = wb.create_sheet("Output Scaling")
    ols = report.get("output_length_scaling", [])
    if ols:
        headers = ["Max New Tokens", "Input Tokens", "New Tokens",
                    "Time (s)", "TTFT (s)", "Tok/s", "ms/tok",
                    "Peak VRAM (MB)", "CPU RAM (MB)"]
        _write_header(ws3, headers)
        fmt = {4: NUM_FMT_2, 5: NUM_FMT_4, 6: NUM_FMT_2, 7: NUM_FMT_2,
               8: NUM_FMT_INT, 9: NUM_FMT_INT}
        for i, o in enumerate(ols, 2):
            _write_row(ws3, i, [
                o.get("max_new_tokens"), o.get("input_tokens"),
                o.get("new_tokens"), o.get("total_time_s"),
                o.get("ttft_s"), o.get("tokens_per_second"),
                o.get("ms_per_token"), o.get("peak_gpu_vram_mb"),
                o.get("cpu_ram_mb"),
            ], fmt)
    _auto_width(ws3)

    # ── Input Scaling sheet ────────────────────────────────────────
    ws4 = wb.create_sheet("Input Scaling")
    ils = report.get("input_length_scaling", [])
    if ils:
        headers = ["Input Chars", "Input Tokens", "New Tokens",
                    "Time (s)", "TTFT (s)", "Tok/s", "ms/tok",
                    "Peak VRAM (MB)", "CPU RAM (MB)"]
        _write_header(ws4, headers)
        fmt = {4: NUM_FMT_2, 5: NUM_FMT_4, 6: NUM_FMT_2, 7: NUM_FMT_2,
               8: NUM_FMT_INT, 9: NUM_FMT_INT}
        for i, o in enumerate(ils, 2):
            _write_row(ws4, i, [
                o.get("input_chars"), o.get("input_tokens"),
                o.get("new_tokens"), o.get("total_time_s"),
                o.get("ttft_s"), o.get("tokens_per_second"),
                o.get("ms_per_token"), o.get("peak_gpu_vram_mb"),
                o.get("cpu_ram_mb"),
            ], fmt)
    _auto_width(ws4)

    # ── Summary sheet ──────────────────────────────────────────────
    ws5 = wb.create_sheet("Summary")
    sm = report.get("summary", {})
    kv5 = [
        ("Avg Tokens/s", sm.get("avg_tokens_per_second", "")),
        ("Avg ms/token", sm.get("avg_ms_per_token", "")),
        ("Peak GPU VRAM (MB)", sm.get("peak_gpu_vram_mb", "")),
        ("Peak CPU RAM (MB)", sm.get("peak_cpu_ram_mb", "")),
        ("Avg TTFT (s)", sm.get("avg_ttft_s", "")),
    ]
    _kv_sheet(ws5, "Summary", kv5)

    _save_wb(wb, out_path)


# ═══════════════════════════════════════════════════════════════════
# 3.  CUMULATIVE WORKBOOK
# ═══════════════════════════════════════════════════════════════════

def write_cumulative_excel(training_reports: List[Dict],
                            inference_reports: List[Dict],
                            out_path: str):
    """Single .xlsx with comparison tables across all profiled runs."""
    wb = Workbook()

    # ── Training Summary ───────────────────────────────────────────
    ws = wb.active
    ws.title = "Training Summary"
    if training_reports:
        headers = [
            "Run Name", "Mode", "Model", "Dtype", "Epochs", "Max Steps",
            "Eff. Batch", "LR", "Final Loss",
            "Samples/s", "Tokens/s", "Steps/s",
            "Total Train (min)", "Peak VRAM (MB)",
            "Train Samples", "Train Tokens",
        ]
        _write_header(ws, headers)
        fmt = {
            8: NUM_FMT_SCI, 9: NUM_FMT_4,
            10: NUM_FMT_2, 11: NUM_FMT_2, 12: NUM_FMT_2,
            13: NUM_FMT_2, 14: NUM_FMT_INT,
            15: NUM_FMT_INT, 16: NUM_FMT_INT,
        }
        for i, r in enumerate(training_reports, 2):
            cfg = r.get("config", {})
            th = r.get("throughput", {})
            mem = r.get("memory", {})
            ti = r.get("timing", {})
            ds = r.get("dataset", {})
            fl = r.get("training_loss", {}).get("final_loss")
            _write_row(ws, i, [
                r.get("run_name"), r.get("mode"), cfg.get("model_name"),
                cfg.get("dtype"), cfg.get("num_epochs"), cfg.get("max_steps"),
                cfg.get("effective_batch"), cfg.get("lr"), fl,
                th.get("samples_per_second"), th.get("tokens_per_second"),
                th.get("steps_per_second"),
                ti.get("total_train_min"), mem.get("peak_gpu_vram_mb"),
                ds.get("train_samples"), ds.get("train_tokens"),
            ], fmt)
    _auto_width(ws)

    # ── Inference Summary ──────────────────────────────────────────
    ws2 = wb.create_sheet("Inference Summary")
    if inference_reports:
        headers = [
            "Model", "Device", "Load Time (s)", "VRAM After Load (MB)",
            "Avg Tok/s", "Avg ms/tok", "Avg TTFT (s)",
            "Peak VRAM (MB)", "Peak CPU RAM (MB)",
        ]
        _write_header(ws2, headers)
        fmt = {
            3: NUM_FMT_2, 4: NUM_FMT_INT,
            5: NUM_FMT_2, 6: NUM_FMT_2, 7: NUM_FMT_4,
            8: NUM_FMT_INT, 9: NUM_FMT_INT,
        }
        for i, r in enumerate(inference_reports, 2):
            ml = r.get("model_load", {})
            sm = r.get("summary", {})
            _write_row(ws2, i, [
                r.get("model_name"), r.get("device"),
                ml.get("load_time_s"), ml.get("gpu_vram_after_load_mb"),
                sm.get("avg_tokens_per_second"), sm.get("avg_ms_per_token"),
                sm.get("avg_ttft_s"), sm.get("peak_gpu_vram_mb"),
                sm.get("peak_cpu_ram_mb"),
            ], fmt)
    _auto_width(ws2)

    # ── GPU-vs-CPU sheet ───────────────────────────────────────────
    gpu = [r for r in inference_reports if r.get("device") == "cuda"]
    cpu = [r for r in inference_reports if r.get("device") == "cpu"]
    gpu_map = {r.get("model_name"): r for r in gpu}
    cpu_map = {r.get("model_name"): r for r in cpu}
    common = sorted(set(gpu_map) & set(cpu_map))

    if common:
        ws3 = wb.create_sheet("GPU vs CPU")
        headers = [
            "Model", "GPU Tok/s", "CPU Tok/s", "Speedup",
            "GPU VRAM (MB)", "CPU RAM (MB)",
        ]
        _write_header(ws3, headers)
        fmt = {2: NUM_FMT_2, 3: NUM_FMT_2, 4: "0.0×",
               5: NUM_FMT_INT, 6: NUM_FMT_INT}
        for i, name in enumerate(common, 2):
            gr = gpu_map[name]
            cr = cpu_map[name]
            g_tps = gr.get("summary", {}).get("avg_tokens_per_second", 0)
            c_tps = cr.get("summary", {}).get("avg_tokens_per_second", 0)
            speedup = g_tps / c_tps if c_tps else 0
            _write_row(ws3, i, [
                name, g_tps, c_tps, speedup,
                gr.get("summary", {}).get("peak_gpu_vram_mb", ""),
                cr.get("summary", {}).get("peak_cpu_ram_mb", ""),
            ], fmt)
        _auto_width(ws3)

    _save_wb(wb, out_path)
