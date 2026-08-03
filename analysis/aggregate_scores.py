#!/usr/bin/env python3
"""Aggregate per-model GPT-5.1 judged workbooks into the paper's master results.

Implements the pre-registered statistical protocol in docs/STUDY_DESIGN.md:
closed-book (P1-P5) vs open-book/RAG (P6-P8) settings, a fixed within-/matched
contrast family, paired Wilcoxon + bootstrap CIs + Cohen's d_z, Benjamini-Hochberg
FDR pooled across the primary family, a practical-significance rule, and
category/difficulty stratified re-runs (statistics only -- no model/judge re-runs).

Reads NetBench-LLM/outputs/by_model/<model>/evaluations/judged/*_by_gpt-5.1.xlsx
and emits, under analysis/outputs/:
  scores_long.csv / scores_overall.csv   per-question scores (regenerable, gitignored)
  summary_by_variant.csv                  per-variant means + n
  completeness_grid.csv                   base_model x P1..P8 presence
  significance.csv                         primary family: Wilcoxon, CI, d_z, BH, practical
  significance_cross_setting.csv           exploratory information-access contrasts
  significance_stratified.csv              primary contrasts within each category/difficulty
  MASTER_RESULTS.md                        human-readable spine

Two overall metrics are emitted side by side (see docs/JUDGE_OVERALL_AUDIT.md):
the judge-reported `overall` (used by paper/, the IEEE baseline) drives the
files above; the deterministic rubric aggregate `overall_formula`
(= 0.4*correctness + 0.3*completeness + 0.2*clarity + 0.1*conciseness, the
EACL paper's primary metric) is an added column in the score/summary CSVs and
drives the parallel significance_formula*.csv files, produced by the identical
procedure. judge_overall_deviation.csv quantifies the judge's arithmetic drift.
"""
from __future__ import annotations
import argparse
import glob
import os

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import wilcoxon, false_discovery_control

DIMENSIONS = ["correctness", "completeness", "clarity", "conciseness", "overall"]

SIZE_GROUP = {
    "llama-3.2-1b": "small", "gemma-3-1b": "small", "qwen3.5-2b": "small",
    "llama-3.1-8b": "medium", "qwen3.5-9b": "medium", "gemma-3-12b": "medium",
    "gemma-3-27b": "large", "qwen3.5-27b": "large",
    "gpt-4o": "api-baseline", "gemini-2.5-pro": "api-baseline",
    "claude-sonnet-4-6": "api-baseline",
}

# --- pre-registered thresholds (see docs/STUDY_DESIGN.md) ---
FDR_Q = 0.05
PRACTICAL_MIN_DELTA = 0.25     # |mean overall difference| on the 1-5 scale
DZ_MIN = 0.2                   # paired Cohen's d_z
MIN_N = 10                     # skip contrasts/strata with fewer paired questions
BOOTSTRAP_N = 10000
SEED = 1234
THINK_PROVISIONAL_PCT = 10.0   # variant flagged provisional if >this% of answers are reasoning-trace-only

# Primary contrast family (a, b, group, description). Positive diff => b > a.
PRIMARY_CONTRASTS = [
    ("P1", "P3", "closed_book", "Instruct vs full SFT"),
    ("P1", "P4", "closed_book", "Instruct vs full CPT+SFT"),
    ("P3", "P4", "closed_book", "full SFT vs full CPT+SFT"),
    ("P2", "P5", "closed_book", "LoRA SFT vs LoRA CPT+SFT"),
    ("P1", "P5", "closed_book", "Instruct vs LoRA CPT+SFT"),
    ("P6", "P7", "open_book", "RAG-Instruct vs RAG-CPT+SFT"),
    ("P6", "P8", "open_book", "RAG-Instruct vs RAG-LoRA-CPT+SFT"),
    ("P1", "P6", "rag_effect", "add RAG to Instruct (matched model)"),
    ("P4", "P7", "rag_effect", "add RAG to full CPT+SFT (matched model)"),
    ("P5", "P8", "rag_effect", "add RAG to LoRA CPT+SFT (matched model)"),
]
# Exploratory, unmatched cross-setting (information-access, NOT model quality).
EXPLORATORY_CROSS = [
    ("P4", "P6", "best closed-adapted (CPT+SFT) vs Instruct+RAG"),
    ("P5", "P6", "best closed-adapted (LoRA CPT+SFT) vs Instruct+RAG"),
]


def setting_of(p_id: str) -> str:
    if p_id in {"P6", "P7", "P8"}:
        return "open_book"
    return "closed_book"   # P1-P5 and API baselines are closed-book


def classify_variant(run_name: str) -> tuple[str, str]:
    n = run_name.lower()
    if n in {m.lower() for m in SIZE_GROUP if SIZE_GROUP[m] == "api-baseline"}:
        return "BASE-API", "API baseline (no adaptation)"
    rag = n.startswith("rag-") or "rag-" in n
    has_cpt, has_lora, has_sft = "cpt" in n, "lora" in n, "sft" in n
    if not has_sft and not has_cpt:
        base = ("P1", "Original Instruct")
    elif has_lora and has_cpt and has_sft:
        base = ("P5", "LoRA CPT + LoRA SFT")
    elif has_lora and has_sft:
        base = ("P2", "LoRA SFT")
    elif has_cpt and has_sft:
        base = ("P4", "Full CPT + Full SFT")
    elif has_sft:
        base = ("P3", "Full SFT")
    else:
        base = ("P?", f"unclassified: {run_name}")
    if rag:
        return {"P1": ("P6", "Instruct + RAG"), "P4": ("P7", "Full CPT+SFT + RAG"),
                "P5": ("P8", "LoRA CPT+SFT + RAG")}.get(base[0], (base[0] + "+RAG", base[1] + " + RAG"))
    return base


def read_judged(path: str) -> pd.DataFrame:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    meta = {k: v for k, v in wb["Metadata"].iter_rows(values_only=True)}
    rows = wb["Judged"].iter_rows(values_only=True)
    header = [str(h) for h in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    out = []
    n_ans = only_think = 0
    for r in rows:
        if r[idx["id"]] is None:
            continue
        ans = str(r[idx["model_answer"]]).lower()
        n_ans += 1
        # "think-only": a reasoning trace with no (or trivial) final answer after </think>
        if "<think>" in ans and ("</think>" not in ans or len(ans.split("</think>")[-1].strip()) < 15):
            only_think += 1
        rec = {"question_id": r[idx["id"]], "category": r[idx["category"]],
               "difficulty": r[idx["difficulty"]]}
        for d in DIMENSIONS:
            v = r[idx[d]]
            rec[d] = float(v) if v not in (None, "") else np.nan
        out.append(rec)
    df = pd.DataFrame(out)
    df["variant_run"] = str(meta.get("model_name", "")).strip()
    df["pct_think_only"] = round(100 * only_think / max(n_ans, 1), 1)
    df["provisional"] = df["pct_think_only"] > THINK_PROVISIONAL_PCT
    return df


def cohen_dz(diff: np.ndarray) -> float:
    sd = diff.std(ddof=1)
    return float(diff.mean() / sd) if sd > 0 else 0.0


def bootstrap_ci(diff: np.ndarray) -> tuple[float, float]:
    rng = np.random.default_rng(SEED)
    means = diff[rng.integers(0, len(diff), size=(BOOTSTRAP_N, len(diff)))].mean(axis=1)
    return float(np.percentile(means, 2.5)), float(np.percentile(means, 97.5))


def paired_test(sa: pd.Series, sb: pd.Series) -> dict | None:
    pair = pd.concat([sa, sb], axis=1).dropna()
    if len(pair) < MIN_N:
        return None
    da, db = pair.iloc[:, 0].to_numpy(), pair.iloc[:, 1].to_numpy()
    diff = db - da
    p = 1.0 if np.allclose(diff, 0) else float(wilcoxon(db, da, zero_method="wilcox").pvalue)
    lo, hi = bootstrap_ci(diff)
    return dict(n=len(pair), mean_a=round(da.mean(), 3), mean_b=round(db.mean(), 3),
                mean_diff=round(float(diff.mean()), 3), median_diff=round(float(np.median(diff)), 3),
                ci95_lo=round(lo, 3), ci95_hi=round(hi, 3), p_raw=p, cohens_dz=round(cohen_dz(diff), 3))


def run_contrasts(allq: pd.DataFrame, models, contrasts, subset=None) -> list[dict]:
    rows = []
    for model in models:
        sub = allq[allq.base_model == model]
        if subset is not None:
            sub = sub[subset(sub)]
        prov = sub.groupby("p_id")["provisional"].max()
        piv = sub.pivot_table(index="question_id", columns="p_id", values="overall")
        for a, b, *rest in contrasts:
            if a not in piv or b not in piv:
                continue
            res = paired_test(piv[a], piv[b])
            if res is None:
                continue
            row = {"base_model": model, "size_group": SIZE_GROUP.get(model),
                   "contrast": f"{a} vs {b}",
                   "provisional": bool(prov.get(a, False) or prov.get(b, False))}
            if len(rest) == 2:
                row["group"], row["description"] = rest
            else:
                row["description"] = rest[0]
            rows.append({**row, **res})
    return rows


def add_practical(df: pd.DataFrame) -> pd.DataFrame:
    df["bh_sig"] = df["p_bh"] <= FDR_Q
    df["practically_sig"] = (df["mean_diff"].abs() >= PRACTICAL_MIN_DELTA) & \
                            (df["cohens_dz"].abs() >= DZ_MIN) & df["bh_sig"]
    return df


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--by-model", default="NetBench-LLM/outputs/by_model")
    ap.add_argument("--out", default="analysis/outputs")
    ap.add_argument("--excluded", default="analysis/excluded_items.csv",
                    help="CSV with an 'id' column of audited-out question IDs (excluded from scoring)")
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)

    frames = []
    for path in sorted(glob.glob(os.path.join(args.by_model, "*/evaluations/judged/*_by_gpt-5.1.xlsx"))):
        model_dir = path.split("/by_model/")[1].split("/")[0]
        df = read_judged(path)
        p_id, label = classify_variant(df["variant_run"].iloc[0])
        df["base_model"], df["size_group"] = model_dir, SIZE_GROUP.get(model_dir, "unknown")
        df["p_id"], df["variant_label"], df["setting"] = p_id, label, setting_of(p_id)
        frames.append(df)
    allq = pd.concat(frames, ignore_index=True)
    # Deterministic rubric aggregate (EACL primary metric; docs/JUDGE_OVERALL_AUDIT.md)
    allq["overall_formula"] = (0.4 * allq["correctness"] + 0.3 * allq["completeness"]
                               + 0.2 * allq["clarity"] + 0.1 * allq["conciseness"])
    print(f"  loaded {allq['variant_run'].nunique()} variants, {allq['base_model'].nunique()} models, "
          f"{allq['question_id'].nunique()} questions")

    # Drop items the expert audit flagged as flawed (wrong reference or dangling
    # source dependency). Same items removed for every model, so relative/paired
    # contrasts are unaffected; see analysis/excluded_items.csv for the list+reasons.
    if args.excluded and os.path.exists(args.excluded):
        excluded = set(pd.read_csv(args.excluded)["id"].astype(str).str.strip())
        before = allq["question_id"].nunique()
        allq = allq[~allq["question_id"].isin(excluded)].copy()
        print(f"  excluded {before - allq['question_id'].nunique()} audited items "
              f"({len(excluded)} listed) -> {allq['question_id'].nunique()} scored")

    base_cols = ["base_model", "size_group", "setting", "p_id", "variant_label", "variant_run",
                 "question_id", "category", "difficulty"]
    metrics = DIMENSIONS + ["overall_formula"]
    allq[base_cols + metrics].to_csv(os.path.join(args.out, "scores_overall.csv"), index=False)
    allq.melt(id_vars=base_cols, value_vars=metrics, var_name="dimension", value_name="score") \
        .to_csv(os.path.join(args.out, "scores_long.csv"), index=False)

    gkeys = ["size_group", "base_model", "setting", "p_id", "variant_label", "variant_run"]
    summ = allq.groupby(gkeys)[metrics].mean().round(3)
    summ["n"] = allq.groupby(gkeys).size()
    summ = summ.join(allq.groupby(gkeys)[["provisional", "pct_think_only"]].max())
    summ.reset_index().sort_values(["size_group", "base_model", "p_id"]).to_csv(
        os.path.join(args.out, "summary_by_variant.csv"), index=False)

    ps = [f"P{i}" for i in range(1, 9)]
    grid = pd.DataFrame(index=[m for m in SIZE_GROUP if SIZE_GROUP[m] != "api-baseline"], columns=ps)
    for m in grid.index:
        present = set(allq[allq.base_model == m]["p_id"])
        grid.loc[m] = ["yes" if p in present else "—" for p in ps]
    grid.to_csv(os.path.join(args.out, "completeness_grid.csv"))

    trained = [m for m in SIZE_GROUP if SIZE_GROUP[m] != "api-baseline"]

    # Judge-reported metric (paper/), then the deterministic rubric aggregate
    # (paper_eacl/) through the identical procedure.
    prim, cross, strat = emit_significance(allq, trained, args.out)
    allq_f = allq.copy()
    allq_f["overall"] = allq_f["overall_formula"]
    emit_significance(allq_f, trained, args.out, suffix="_formula")

    emit_deviation_audit(allq, args.out)

    write_master_md(args.out, allq, grid, prim, cross, strat)
    print(f"\nWrote artifacts to {args.out}/  "
          f"(primary contrasts: {len(prim)}, stratified: {len(strat)}, cross-setting: {len(cross)}; "
          f"judge + formula metrics)")


def emit_significance(allq, trained, out, suffix=""):
    # --- primary family: BH pooled across all (model x contrast) ---
    prim = pd.DataFrame(run_contrasts(allq, trained, PRIMARY_CONTRASTS))
    ok = ~prim["provisional"]                      # BH only over non-provisional contrasts
    prim["p_bh"] = np.nan
    prim.loc[ok, "p_bh"] = false_discovery_control(prim.loc[ok, "p_raw"].to_numpy(), method="bh")
    prim["p_bh"] = prim["p_bh"].round(5)
    prim["p_raw"] = prim["p_raw"].round(5)
    prim = add_practical(prim)                      # NaN p_bh -> bh_sig/practically_sig = False
    prim = prim[["base_model", "size_group", "group", "contrast", "description", "n",
                 "mean_a", "mean_b", "mean_diff", "median_diff", "ci95_lo", "ci95_hi",
                 "p_raw", "p_bh", "bh_sig", "cohens_dz", "practically_sig", "provisional"]]
    prim.to_csv(os.path.join(out, f"significance{suffix}.csv"), index=False)

    # --- exploratory cross-setting (no BH; flagged information-access) ---
    cross = pd.DataFrame(run_contrasts(allq, trained, EXPLORATORY_CROSS))
    cross["p_raw"] = cross["p_raw"].round(5)
    cross["note"] = "information-access contrast; NOT a model-quality claim"
    cross.to_csv(os.path.join(out, f"significance{suffix}_cross_setting.csv"), index=False)

    # --- stratified: primary contrasts within each category / difficulty (BH within each stratum) ---
    strat_rows = []
    for stype, col in [("category", "category"), ("difficulty", "difficulty")]:
        for sval in sorted(allq[col].dropna().unique()):
            rs = run_contrasts(allq, trained, PRIMARY_CONTRASTS, subset=lambda d, c=col, v=sval: d[c] == v)
            for r in rs:
                r["stratum_type"], r["stratum"] = stype, sval
            strat_rows += rs
    strat = pd.DataFrame(strat_rows)
    if not strat.empty:
        strat = strat[~strat["provisional"]].copy()   # exclude think-corrupted contrasts
        strat["p_bh"] = strat.groupby(["stratum_type", "stratum"])["p_raw"].transform(
            lambda s: false_discovery_control(s.to_numpy(), method="bh"))
        strat["p_raw"], strat["p_bh"] = strat["p_raw"].round(5), strat["p_bh"].round(5)
        strat = add_practical(strat)
        strat["exploratory"] = True
        strat = strat[["stratum_type", "stratum", "base_model", "group", "contrast", "n",
                       "mean_diff", "ci95_lo", "ci95_hi", "p_raw", "p_bh", "bh_sig",
                       "cohens_dz", "practically_sig", "exploratory"]]
        strat.to_csv(os.path.join(out, f"significance{suffix}_stratified.csv"), index=False)
    return prim, cross, strat


def emit_deviation_audit(allq, out):
    """judge_overall_deviation.csv: how the judge-reported overall deviates from
    the instructed weighted formula, by slice, with the drift-toward-unweighted-
    mean diagnostics (docs/JUDGE_OVERALL_AUDIT.md)."""
    d = allq.dropna(subset=DIMENSIONS).copy()
    d["dev"] = d["overall"] - d["overall_formula"]
    unw = d[["correctness", "completeness", "clarity", "conciseness"]].mean(axis=1)
    d["toward"] = unw - d["overall_formula"]
    rows = []

    def slice_stats(name, g):
        nz = g[g["dev"].abs() > 0.005]
        rows.append(dict(
            slice=name, n=len(g),
            mad=round(g["dev"].abs().mean(), 4),
            signed_bias=round(g["dev"].mean(), 4),
            exact_pct=round((g["dev"].abs() < 0.005).mean() * 100, 1),
            gt_0_25_pct=round((g["dev"].abs() > 0.25).mean() * 100, 1),
            max_abs=round(g["dev"].abs().max(), 2),
            drift_corr=round(float(np.corrcoef(g["dev"], g["toward"])[0, 1]), 3)
                       if g["dev"].std() > 0 else np.nan,
            drift_sign_share=round(float((np.sign(nz["dev"]) == np.sign(nz["toward"])).mean()), 3)
                             if len(nz) else np.nan))

    slice_stats("all", d)
    slice_stats("closed_book_local", d[(d["setting"] == "closed_book") & (d["p_id"] != "BASE-API")])
    slice_stats("open_book", d[d["setting"] == "open_book"])
    slice_stats("api_baseline", d[d["p_id"] == "BASE-API"])
    for p in sorted(d["p_id"].unique()):
        slice_stats(p, d[d["p_id"] == p])
    for m in sorted(d["base_model"].unique()):
        slice_stats(m, d[d["base_model"] == m])
    pd.DataFrame(rows).to_csv(os.path.join(out, "judge_overall_deviation.csv"), index=False)


def write_master_md(out, allq, grid, prim, cross, strat):
    L = ["# NetBench — Master Results (auto-generated)\n",
         "Source: consolidated GPT-5.1 judged workbooks under "
         "`NetBench-LLM/outputs/by_model/`. Regenerate with "
         "`analysis/.venv/bin/python analysis/aggregate_scores.py`. "
         "Protocol: see `docs/STUDY_DESIGN.md` (statistical analysis protocol).\n",
         f"Coverage: **{allq['base_model'].nunique()} models**, "
         f"**{allq['variant_run'].nunique()} model-variants**, "
         f"**{allq['question_id'].nunique()} questions** (judge `gpt-5.1`, HPN v5.0). "
         f"Settings: **closed-book** = P1–P5, **open-book (RAG)** = P6–P8.\n",
         "Metric note: this report and `significance*.csv` use the judge-reported "
         "`overall` (IEEE `paper/`). The EACL paper uses the deterministic "
         "`overall_formula` (`significance_formula*.csv`); see "
         "`docs/JUDGE_OVERALL_AUDIT.md`.\n"]

    prov_variants = allq[allq["provisional"]].groupby(["base_model", "p_id"])["pct_think_only"].max()
    if len(prov_variants):
        L.append("\n## ⚠ Data quality — provisional cells\n")
        L.append(f"{len(prov_variants)} variant(s) are **think-corrupted** "
                 f"(>{THINK_PROVISIONAL_PCT:.0f}% of answers are reasoning-trace-only with no final answer); "
                 "they are **excluded from primary significance** and any cell/row touching them is "
                 "**provisional** pending regeneration:\n")
        for (m, p), pct in prov_variants.items():
            L.append(f"- `{m}` {p} — {pct:.0f}% think-only")
        L.append("")
    L.append("\n## Overall score by model × variant (P-taxonomy)\n")
    L.append("Closed-book: P1–P5 · Open-book/RAG: P6–P8 · BASE-API = frontier baselines (closed-book). "
             "Cells for the provisional variants listed above are not yet reliable.\n")
    piv = allq.pivot_table(index="base_model", columns="p_id", values="overall", aggfunc="mean").round(2)
    row_order = [m for m in SIZE_GROUP if m in piv.index]
    piv = piv.reindex(row_order)
    piv.insert(0, "size", [SIZE_GROUP.get(m, "") for m in piv.index])
    piv.index.name = "model"
    order = ["size"] + [p for p in [f"P{i}" for i in range(1, 9)] + ["BASE-API"] if p in piv.columns]
    L.append(piv[order].to_markdown())

    L.append("\n\n## Completeness grid (which variants exist)\n")
    L.append(grid.to_markdown())

    L.append("\n\n## Primary significance (pre-registered family, BH-FDR pooled, q=0.05)\n")
    L.append("Positive `mean_diff` ⇒ second variant higher. `p_bh` = Benjamini–Hochberg adjusted; "
             "**`practically_sig`** requires |mean_diff| ≥ 0.25 *and* |d_z| ≥ 0.2 *and* BH-significant. "
             "Groups: closed_book / open_book (within-setting adaptation) and rag_effect (matched model ± retrieval).\n")
    show = prim[["base_model", "group", "contrast", "description", "n", "mean_diff",
                 "ci95_lo", "ci95_hi", "p_bh", "cohens_dz", "practically_sig", "provisional"]]
    L.append(show.to_markdown(index=False))
    n_prov = int(prim["provisional"].sum())
    n_ps = int(prim["practically_sig"].sum())
    L.append(f"\n\n**Practically-significant primary contrasts: {n_ps} / {len(prim)}** "
             f"({n_prov} excluded as provisional/think-corrupted, not BH-tested). "
             "Many statistically-significant gaps still fail the practical bar — see the discussion of "
             "what counts as a meaningful improvement.\n")

    L.append("\n## Cross-setting (exploratory — information access, NOT model quality)\n")
    L.append("RAG variants retrieve source passages the closed-book models never see; "
             "these contrasts are reported only as an information-access comparison.\n")
    if not cross.empty:
        L.append(cross[["base_model", "contrast", "description", "mean_diff", "ci95_lo", "ci95_hi", "p_raw"]].to_markdown(index=False))

    L.append("\n\n## Stratified highlights (exploratory)\n")
    if strat is not None and not strat.empty:
        flips = []
        pm = prim.set_index(["base_model", "contrast"])["mean_diff"]
        for (m, c), g in strat.groupby(["base_model", "contrast"]):
            if (m, c) in pm.index:
                full = pm.loc[(m, c)]
                opp = g[np.sign(g["mean_diff"]) != np.sign(full)]
                for _, r in opp.iterrows():
                    flips.append(f"- `{m}` {c}: overall {full:+.2f} but {r['stratum']} {r['mean_diff']:+.2f}")
        L.append(f"Per-stratum tests live in `significance_stratified.csv` "
                 f"({len(strat)} rows; strata with n<{MIN_N} are skipped). ")
        if flips:
            L.append(f"Sign-flip strata (effect reverses vs the full-set result), {len(flips)} total — first 12:\n")
            L.append("\n".join(flips[:12]))
        else:
            L.append("No contrast reverses sign in any stratum.")
    open(os.path.join(out, "MASTER_RESULTS.md"), "w").write("\n".join(L) + "\n")


if __name__ == "__main__":
    main()
