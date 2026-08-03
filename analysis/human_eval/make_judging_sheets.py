#!/usr/bin/env python3
"""Build the blinded Excel workbook for the human-judge alignment study.

Design (see README.md in this directory):
- 5 variants spanning the score range and embedding headline contrasts:
  qwen3.5-2b P1 / P4 / P6 (shared-baseline FT and RAG effects),
  qwen3.5-9b P6 (best open RAG, frontier-parity claim), gpt-4o (frontier).
- 25 questions: 19 of the first 25 benchmark items (minus the one
  audit-excluded item and 5 drops from over-represented cells), plus 2 fixed
  adds covering the missing tenth category and the calculation-item quota,
  plus 4 seeded-random fills to reach 8 easy / 10 medium / 7 hard.
- One blinded Scoring sheet: per question, the 5 responses appear in
  seeded-random order under codes Q<nn>-a..e. The identity mapping goes to
  sampling_key.csv, which the rater must not open until scoring is done.

The rater sees exactly what the GPT-5.1 judge saw: question, reference
answer, model answer (rubric text below is verbatim from
NetBench-LLM/evaluation/judge_responses.py).
"""

import csv
import json
import math
import random
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
BENCHMARK = ROOT / "NetBench-LLM/data/prompts/hpn_benchmark_v5.0.jsonl"
BY_MODEL = ROOT / "NetBench-LLM/outputs/by_model"
OUTDIR = Path(__file__).resolve().parent / "outputs"

SEED = 42

# (base_model, p_id, setting, variant_run, judged file relative to by_model/)
VARIANTS = [
    ("qwen3.5-2b", "P1", "closed_book", "Qwen3.5-2B",
     "qwen3.5-2b/evaluations/judged/hpn_judged_Qwen3.5-2B_by_gpt-5.1.xlsx"),
    ("qwen3.5-2b", "P4", "closed_book", "Qwen3.5-2B-cpt-full-sft",
     "qwen3.5-2b/evaluations/judged/hpn_judged_Qwen3.5-2B-cpt-full-sft_by_gpt-5.1.xlsx"),
    ("qwen3.5-2b", "P6", "open_book", "RAG-Qwen3.5-2B",
     "qwen3.5-2b/evaluations/judged/hpn_judged_RAG-Qwen3.5-2B_by_gpt-5.1.xlsx"),
    ("qwen3.5-9b", "P6", "open_book", "RAG-Qwen3.5-9B",
     "qwen3.5-9b/evaluations/judged/hpn_judged_RAG-Qwen3.5-9B_by_gpt-5.1.xlsx"),
    ("gpt-4o", "BASE-API", "closed_book", "gpt-4o",
     "gpt-4o/evaluations/judged/hpn_judged_gpt-4o_by_gpt-5.1.xlsx"),
]

# Drops from the first-25 pool (user has studied these items in depth, so we
# keep most): trim over-represented cells to make room for stratification.
# 3 of 7 Bottleneck Diagnosis, 1 of 3 BDP medium, 1 of 2 Transfer easy.
DROP_FROM_FIRST25 = {
    "NB-HPN-0733-q1",  # Bottleneck Diagnosis, easy   (keeps NB-HPN-2032-q1)
    "NB-HPN-0850-q2",  # Bottleneck Diagnosis, hard   (keeps NB-HPN-1278-q1)
    "NB-HPN-1376-q1",  # Bottleneck Diagnosis, medium (keeps 1803-q3, 0555-q1)
    "NB-HPN-0178-q1",  # BDP-Based Reasoning, medium  (keeps NB-HPN-0176-q1)
    "NB-HPN-1135-q0",  # Transfer Parameters, easy    (keeps NB-HPN-1629-q1)
}
# Fixed adds: the tenth category absent from the first 25 (Dataset
# Partitioning and Mixed Workloads); 2032-q3 also fills the calc-item quota.
FIXED_ADDS = ["NB-HPN-1676-q1", "NB-HPN-2032-q3"]
# Remaining slots drawn uniformly at random (SEED) within difficulty strata.
RANDOM_FILL = {"easy": 2, "medium": 1, "hard": 1}

# Verbatim from NetBench-LLM/evaluation/judge_responses.py (JUDGE_SYSTEM_PROMPT
# scoring principles and the JUDGE_USER_TEMPLATE dimension anchors).
SCORING_PRINCIPLES = [
    "Treat the QUESTION as the ultimate source of truth for what is being asked.",
    "Treat the REFERENCE ANSWER as a high-quality exemplar, NOT the only valid "
    "answer. Give full credit to responses that are technically correct and "
    "address the question, even if worded differently from the reference.",
    "If the model answer contains additional content beyond the reference that "
    "is factually correct, treat it as neutral for correctness but potentially "
    "negative for conciseness.",
    "Be strict but fair. Focus on technical accuracy. Partial credit is allowed.",
]
RUBRIC = [
    ("Correctness", "Is the technical content accurate relative to the reference?", [
        "1 = Fundamentally wrong or missing key facts",
        "2 = Mostly wrong, a few correct fragments",
        "3 = Partially correct, some errors or omissions",
        "4 = Mostly correct, minor inaccuracies",
        "5 = Fully correct, matches reference in all key technical claims"]),
    ("Completeness", "Does the answer cover all important concepts from the reference?", [
        "1 = Covers <20% of key concepts",
        "2 = Covers 20-40%",
        "3 = Covers 40-60%",
        "4 = Covers 60-80%",
        "5 = Covers >80% of key concepts"]),
    ("Clarity", "Is the answer well-organized and easy to follow?", [
        "1 = Incoherent or very hard to follow",
        "2 = Mostly confusing with occasional clear points",
        "3 = Understandable but poorly structured",
        "4 = Well-structured with minor clarity issues",
        "5 = Exceptionally clear and well-organized"]),
    ("Conciseness", "Is the answer appropriately focused without unnecessary padding?", [
        "1 = Extremely padded, repetitive, or off-topic",
        "2 = Significant padding or irrelevant tangents",
        "3 = Acceptable length with some unnecessary content",
        "4 = Mostly focused, minimal padding",
        "5 = Perfectly concise - every sentence adds value"]),
]


def load_data():
    bench = [json.loads(l) for l in open(BENCHMARK, encoding="utf-8")]
    meta = {r["id"]: r for r in bench}
    order = [r["id"] for r in bench]
    judged = {}
    for base_model, p_id, setting, run, rel in VARIANTS:
        df = pd.read_excel(BY_MODEL / rel, sheet_name="Judged")
        assert list(df.id) == order, f"question order differs in {rel}"
        judged[run] = df.set_index("id")
    return meta, order, judged


def pick_questions(meta, order, judged):
    ref = judged[VARIANTS[0][3]]
    excluded = set(ref[ref.excluded_from_scoring].index)
    first25 = order[:25]
    kept = [q for q in first25 if q not in excluded and q not in DROP_FROM_FIRST25]
    chosen = kept + FIXED_ADDS
    rng = random.Random(SEED)
    pool = [q for q in order if q not in excluded and q not in first25
            and q not in chosen]
    for diff, k in RANDOM_FILL.items():
        cands = sorted(q for q in pool if meta[q]["difficulty"] == diff)
        chosen += rng.sample(cands, k)
    # present in benchmark order
    chosen = [q for q in order if q in set(chosen)]
    assert len(chosen) == 25
    return chosen


def est_height(text, chars_per_line, line_pt=13.0, min_pt=15.0, max_pt=409.0):
    lines = sum(max(1, math.ceil(len(seg) / chars_per_line))
                for seg in str(text).split("\n"))
    return max(min_pt, min(max_pt, lines * line_pt + 4))


THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def build_instructions(ws):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 112
    rows = [
        ("HPN-QA Human Judging - Instructions", "title"),
        ("", None),
        ("Purpose: score model-generated answers against expert reference "
         "answers, using exactly the rubric the LLM judge used. Your scores "
         "will be compared with the LLM judge's scores to measure "
         "human-judge alignment.", None),
        ("", None),
        ("Blinding: system identities are hidden and the order of the five "
         "responses is randomized independently for every question. Do NOT "
         "open sampling_key.csv until you have finished scoring the entire "
         "sheet.", "bold"),
        ("", None),
        ("Procedure, for each question block on the Scoring sheet:", "head"),
        ("  1. Read the QUESTION row and the REFERENCE ANSWER row.", None),
        ("  2. For each response row (codes Qnn-a ... Qnn-e), read the answer "
         "and fill the four score cells with whole numbers 1-5 using the "
         "anchors below. Score every response on its own merits against the "
         "reference - avoid ranking the five responses against each other.", None),
        ("  3. Do NOT compute an overall score. It is derived later as the "
         "same weighted composite the judge was instructed to use "
         "(0.4*correctness + 0.3*completeness + 0.2*clarity + "
         "0.1*conciseness).", None),
        ("", None),
        ("Scoring principles (verbatim from the judge prompt):", "head"),
        *[(f"  - {p}", None) for p in SCORING_PRINCIPLES],
        ("", None),
    ]
    for name, desc, anchors in RUBRIC:
        rows.append((f"{name} (1-5): {desc}", "head"))
        rows += [(f"  {a}", None) for a in anchors]
        rows.append(("", None))
    rows.append(("When finished: save a copy named "
                 "HPN_human_judging_v1_filled_<your-initials>.xlsx.", "bold"))
    for i, (text, style) in enumerate(rows, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = WRAP
        if style == "title":
            c.font = Font(bold=True, size=14)
        elif style in ("head", "bold"):
            c.font = Font(bold=True)
        ws.row_dimensions[i].height = est_height(text, 105) if text else 8


def build_scoring(ws, meta, chosen, judged, rng):
    header = ["ID", "Content", "Correctness", "Completeness", "Clarity",
              "Conciseness"]
    widths = [10, 100, 12, 13, 9, 12]
    for j, (h, w) in enumerate(zip(header, widths), start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="404040")
        c.alignment = CENTER
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"

    dv = DataValidation(type="list", formula1='"1,2,3,4,5"', allow_blank=True,
                        showErrorMessage=True,
                        error="Enter a whole number from 1 to 5.")
    ws.add_data_validation(dv)

    q_fill = PatternFill("solid", fgColor="DDEBF7")
    ref_fill = PatternFill("solid", fgColor="E2EFDA")
    score_fill = PatternFill("solid", fgColor="FFF2CC")

    key_rows = []
    r = 2
    for i, qid in enumerate(chosen, start=1):
        qno = f"Q{i:02d}"
        question = meta[qid]["question"]
        reference = meta[qid]["reference_answer"]
        for label, text, fill in [(qno, "QUESTION: " + question, q_fill),
                                  ("Ref", "REFERENCE ANSWER: " + reference, ref_fill)]:
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            ws.cell(row=r, column=1).fill = fill
            ws.cell(row=r, column=1).alignment = WRAP
            c = ws.cell(row=r, column=2, value=text)
            c.fill = fill; c.alignment = WRAP
            c.font = Font(bold=(label == qno))
            for j in range(3, 7):
                ws.cell(row=r, column=j).fill = fill
            ws.row_dimensions[r].height = est_height(text, 95)
            r += 1
        shuffled = VARIANTS[:]
        rng.shuffle(shuffled)
        for pos, (base_model, p_id, setting, run, rel) in enumerate(shuffled):
            rid = f"{qno}-{'abcde'[pos]}"
            answer = judged[run].loc[qid, "model_answer"]
            assert isinstance(answer, str) and answer.strip(), f"empty answer {run} {qid}"
            ws.cell(row=r, column=1, value=rid).alignment = WRAP
            c = ws.cell(row=r, column=2, value=answer)
            c.alignment = WRAP
            for j in range(3, 7):
                sc = ws.cell(row=r, column=j)
                sc.fill = score_fill; sc.alignment = CENTER
                dv.add(sc)
            for j in range(1, 7):
                ws.cell(row=r, column=j).border = THIN
            ws.row_dimensions[r].height = est_height(answer, 95)
            key_rows.append({
                "response_id": rid, "question_no": qno, "question_id": qid,
                "base_model": base_model, "p_id": p_id, "setting": setting,
                "variant_run": run, "judged_file": rel,
                "category": meta[qid]["category"],
                "difficulty": meta[qid]["difficulty"],
                "question_type": meta[qid]["question_type"],
                "requires_calculation": meta[qid]["requires_calculation"],
            })
            r += 1
    return key_rows


def main():
    OUTDIR.mkdir(exist_ok=True)
    meta, order, judged = load_data()
    chosen = pick_questions(meta, order, judged)
    rng = random.Random(SEED)

    wb = Workbook()
    build_instructions(wb.active)
    wb.active.title = "Instructions"
    key_rows = build_scoring(wb.create_sheet("Scoring"), meta, chosen, judged, rng)
    wb.save(OUTDIR / "HPN_human_judging_v1.xlsx")

    with open(OUTDIR / "sampling_key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(key_rows[0]))
        w.writeheader(); w.writerows(key_rows)

    diffs = pd.Series([meta[q]["difficulty"] for q in chosen]).value_counts()
    cats = len({meta[q]["category"] for q in chosen})
    calcs = sum(meta[q]["requires_calculation"] for q in chosen)
    from25 = sum(q in order[:25] for q in chosen)
    print(f"25 questions ({from25} from the first 25), "
          f"difficulty {diffs.to_dict()}, {cats} categories, {calcs} calc items; "
          f"{len(key_rows)} responses -> {OUTDIR}")


if __name__ == "__main__":
    main()
