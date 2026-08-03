#!/usr/bin/env python3
"""
RAG Pipeline Benchmark — LLM-as-Judge (Phase 2)
================================================

Reads Phase 1 answer Excel files, sends each (question, reference_answer,
model_answer) triple to a judge LLM (Gemini or OpenAI) for scoring, and
saves judged results to Excel.

Features:
    • Rate limiting with configurable --requests_per_minute (default: 15)
    • Exponential backoff on 429 / ResourceExhausted errors
    • Resume capability — skips already-judged question IDs if the
      output file already exists
    • JSON parse with one retry on malformed responses
    • Excel output with 3 sheets: Metadata, Judged, Summary
    • Supports Gemini models (gemini-*) and OpenAI models (gpt-*, o1-*, o3-*, o4-*)

Output:
    outputs/judged/hpn_judged_{model_name}_{benchmark_tag}_by_{judge_model}.xlsx

Scoring rubric (identical to LLM-Training evaluation for direct comparison):
    Correctness  × 0.40  — Technical accuracy vs. reference
    Completeness × 0.30  — Coverage of key concepts
    Clarity      × 0.20  — Organization and readability
    Conciseness  × 0.10  — Appropriate focus without padding

Usage:
    # Judge with Gemini (default)
    python evaluation/judge_responses.py \\
        --answer_files outputs/answers/hpn_answers_RAG-gpt-4o.xlsx

    # Judge with OpenAI GPT-4o
    python evaluation/judge_responses.py \\
        --answer_files outputs/answers/hpn_answers_RAG-gpt-4o.xlsx \\
        --judge_model gpt-4o \\
        --openai_api_key_file openai_api_key.txt

    # Custom rate limit and judge model
    python evaluation/judge_responses.py \\
        --answer_files outputs/answers/hpn_answers_RAG-gpt-4o.xlsx \\
        --judge_model gemini-2.5-pro \\
        --requests_per_minute 30
"""

import argparse
import json
import os
import re
import sys
import time
import warnings
from datetime import datetime
from pathlib import Path

import yaml

# Add project root to path
ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from google import genai
from google.genai import types as genai_types
from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")


# ═══════════════════════════════════════════════════════════════════════
# Constants
# ═══════════════════════════════════════════════════════════════════════

DEFAULT_CONFIG = str(ROOT / "config.yaml")
DEFAULT_OUTPUT_DIR = str(ROOT / "outputs" / "judged")
DEFAULT_JUDGE_MODEL = "gemini-2.5-flash"
DEFAULT_RPM = 15
MAX_RETRIES = 5
MAX_BACKOFF_SEC = 60

# Prefixes used to detect OpenAI models vs Gemini
_OPENAI_PREFIXES = ("gpt-", "o1-", "o3-", "o4-")


def _is_openai_model(model_name: str) -> bool:
    """Return True if *model_name* is an OpenAI model."""
    return model_name.lower().startswith(_OPENAI_PREFIXES)

# ── Judge prompts (identical to LLM-Training for comparable scoring) ────

JUDGE_SYSTEM_PROMPT = (
    "You are an expert evaluator specializing in "
    "high-performance networking (HPN), HPC data transfer systems, TCP/IP networking, "
    "and computer networks. Your task is to score a model-generated answer "
    "against an expert reference answer.\n\n"
    "Scoring principles:\n"
    "- Treat the QUESTION as the ultimate source of truth for what is being asked.\n"
    "- Treat the REFERENCE ANSWER as a high-quality exemplar, NOT the only valid "
    "answer. Give full credit to responses that are technically correct and address "
    "the question, even if worded differently from the reference.\n"
    "- If the model answer contains additional content beyond the reference that is "
    "factually correct, treat it as neutral for correctness but potentially negative "
    "for conciseness.\n"
    "- Be strict but fair. Focus on technical accuracy. Partial credit is allowed."
)

JUDGE_USER_TEMPLATE = """\
## Question
{question}

## Reference Answer (Expert-Written)
{reference_answer}

## Model Answer
{model_answer}

---
Score the model answer on the following four dimensions using a 1–5 scale:

**Correctness** (1–5): Is the technical content accurate relative to the reference?
  1 = Fundamentally wrong or missing key facts
  2 = Mostly wrong, a few correct fragments
  3 = Partially correct, some errors or omissions
  4 = Mostly correct, minor inaccuracies
  5 = Fully correct, matches reference in all key technical claims

**Completeness** (1–5): Does the answer cover all important concepts from the reference?
  1 = Covers <20% of key concepts
  2 = Covers 20–40%
  3 = Covers 40–60%
  4 = Covers 60–80%
  5 = Covers >80% of key concepts

**Clarity** (1–5): Is the answer well-organized and easy to follow?
  1 = Incoherent or very hard to follow
  2 = Mostly confusing with occasional clear points
  3 = Understandable but poorly structured
  4 = Well-structured with minor clarity issues
  5 = Exceptionally clear and well-organized

**Conciseness** (1–5): Is the answer appropriately focused without unnecessary padding?
  1 = Extremely padded, repetitive, or off-topic
  2 = Significant padding or irrelevant tangents
  3 = Acceptable length with some unnecessary content
  4 = Mostly focused, minimal padding
  5 = Perfectly concise — every sentence adds value

Respond ONLY with a valid JSON object in exactly this format (no other text):
{{
  "justification": "<2-3 sentence explanation of the scores>",
  "correctness": <int 1-5>,
  "completeness": <int 1-5>,
  "clarity": <int 1-5>,
  "conciseness": <int 1-5>,
  "overall": <float, weighted average: correctness*0.4 + completeness*0.3 + clarity*0.2 + conciseness*0.1>
}}"""


# ═══════════════════════════════════════════════════════════════════════
# Excel style constants (identical to LLM-Training Phase 2)
# ═══════════════════════════════════════════════════════════════════════

_HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
_THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
_META_KEY_FONT = Font(name="Calibri", bold=True, size=11)
_SCORE_GOOD_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_SCORE_MED_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_SCORE_BAD_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")


def _style_header_row(ws, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER


def _style_header_row_at(ws, row_num: int, num_cols: int):
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER


def _score_fill(score: float) -> PatternFill:
    if score >= 4.0:
        return _SCORE_GOOD_FILL
    elif score >= 3.0:
        return _SCORE_MED_FILL
    else:
        return _SCORE_BAD_FILL


def _sanitize_for_filename(name: str) -> str:
    return re.sub(r'[^\w\-.]', '_', name)


def _benchmark_tag(benchmark_path: str) -> str:
    """Extract a concise tag from the benchmark filename.

    Examples:
        'hpn_qa_benchmark_v4_general_skills.json' → 'v4_general_skills'
        'hpn_benchmark_v5.0.jsonl'                → 'v5.0'
        'hpn_qa_benchmark.json'                   → ''  (no tag)
    """
    stem = Path(benchmark_path).stem
    for prefix in ("hpn_qa_benchmark", "hpn_benchmark"):
        if stem.startswith(prefix) and len(stem) > len(prefix):
            return stem[len(prefix) + 1:]   # strip prefix + underscore
    return ""


# ═══════════════════════════════════════════════════════════════════════
# API key loading
# ═══════════════════════════════════════════════════════════════════════

def load_api_key(key_file: str | None, cfg: dict | None = None) -> str:
    """Load the Gemini API key.

    Priority:
        1. --gemini_api_key_file CLI flag (pointing to a plain-text file)
        2. paths.gemini_api_key from config.yaml
        3. GEMINI_API_KEY environment variable
    """
    # Try explicit key file first
    if key_file:
        path = Path(key_file)
        if path.exists():
            api_key = path.read_text(encoding="utf-8").strip()
            if api_key:
                print(f"🔑 API key loaded from: {key_file}")
                return api_key
        print(f"⚠️  Key file not found or empty: {key_file}")

    # Try config.yaml path
    if cfg:
        cfg_key_path = cfg.get("paths", {}).get("gemini_api_key", "")
        if cfg_key_path:
            path = Path(cfg_key_path)
            if not path.is_absolute():
                path = ROOT / path
            if path.exists():
                api_key = path.read_text(encoding="utf-8").strip()
                if api_key:
                    print(f"🔑 API key loaded from config path: {path}")
                    return api_key

    # Fall back to env var
    api_key = os.environ.get("GEMINI_API_KEY", "").strip()
    if api_key:
        print("🔑 API key loaded from GEMINI_API_KEY environment variable")
        return api_key

    print("❌ No Gemini API key found.")
    print("   Provide one via:")
    print("     --gemini_api_key_file <path>")
    print("     paths.gemini_api_key in config.yaml")
    print("     GEMINI_API_KEY=<key>  (env var)")
    sys.exit(1)


def load_openai_api_key(key_file: str | None, cfg: dict | None = None) -> str:
    """Load the OpenAI API key.

    Priority:
        1. --openai_api_key_file CLI flag (pointing to a plain-text file)
        2. paths.openai_api_key from config.yaml
        3. OPENAI_API_KEY environment variable
    """
    # Try explicit key file first
    if key_file:
        path = Path(key_file)
        if path.exists():
            api_key = path.read_text(encoding="utf-8").strip()
            if api_key:
                print(f"🔑 OpenAI API key loaded from: {key_file}")
                return api_key
        print(f"⚠️  Key file not found or empty: {key_file}")

    # Try config.yaml path
    if cfg:
        cfg_key_path = cfg.get("paths", {}).get("openai_api_key", "")
        if cfg_key_path:
            path = Path(cfg_key_path)
            if not path.is_absolute():
                path = ROOT / path
            if path.exists():
                api_key = path.read_text(encoding="utf-8").strip()
                if api_key:
                    print(f"🔑 OpenAI API key loaded from config path: {path}")
                    return api_key

    # Fall back to env var
    api_key = os.environ.get("OPENAI_API_KEY", "").strip()
    if api_key:
        print("🔑 OpenAI API key loaded from OPENAI_API_KEY environment variable")
        return api_key

    print("❌ No OpenAI API key found.")
    print("   Provide one via:")
    print("     --openai_api_key_file <path>")
    print("     paths.openai_api_key in config.yaml")
    print("     OPENAI_API_KEY=<key>  (env var)")
    sys.exit(1)


# ═══════════════════════════════════════════════════════════════════════
# Read Phase 1 answer Excel file
# ═══════════════════════════════════════════════════════════════════════

def read_answer_excel(filepath: str) -> tuple[dict, list[dict]]:
    """Read a Phase 1 answer Excel file.

    Returns:
        (metadata_dict, list_of_answer_dicts)
    """
    wb = load_workbook(filepath, read_only=True, data_only=True)

    ws_meta = wb["Metadata"]
    metadata = {}
    for row in ws_meta.iter_rows(min_row=2, max_col=2, values_only=True):
        if row[0] is not None:
            metadata[str(row[0])] = str(row[1]) if row[1] is not None else ""

    ws_ans = wb["Answers"]
    headers = [cell.value for cell in next(ws_ans.iter_rows(min_row=1, max_row=1))]

    answers = []
    for row in ws_ans.iter_rows(min_row=2, values_only=True):
        record = {}
        for col_idx, header in enumerate(headers):
            record[header] = row[col_idx] if col_idx < len(row) else None
        answers.append(record)

    wb.close()
    return metadata, answers


# ═══════════════════════════════════════════════════════════════════════
# Resume support
# ═══════════════════════════════════════════════════════════════════════

def load_existing_judgments(output_path: str) -> dict[str, dict]:
    """If an output file already exists, load judged rows by question ID."""
    if not Path(output_path).exists():
        return {}

    try:
        wb = load_workbook(output_path, read_only=True, data_only=True)
        if "Judged" not in wb.sheetnames:
            wb.close()
            return {}

        ws = wb["Judged"]
        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

        existing = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            record = {}
            for col_idx, header in enumerate(headers):
                record[header] = row[col_idx] if col_idx < len(row) else None
            if record.get("overall") is not None:
                existing[str(record.get("id", ""))] = record

        wb.close()
        print(f"♻️  Found {len(existing)} already-judged questions in {output_path}")
        return existing

    except Exception as e:
        print(f"⚠️  Could not load existing judgments: {e}")
        return {}


# ═══════════════════════════════════════════════════════════════════════
# LLM judge API calls with rate limiting + exponential backoff
# ═══════════════════════════════════════════════════════════════════════

def _loads_lenient(s: str):
    """json.loads, tolerant of invalid backslash escapes (e.g. LaTeX such as
    \\alpha) that judges sometimes emit inside JSON string values."""
    try:
        return json.loads(s)
    except json.JSONDecodeError:
        # Escape any backslash that does not begin a valid JSON escape.
        return json.loads(re.sub(r'\\(?!["\\/bfnrtu])', r"\\\\", s))


def _parse_judge_json(text: str) -> dict | None:
    """Extract and parse the JSON object from the judge response.

    Backslash escapes that are invalid JSON (common when justifications contain
    LaTeX) are repaired before parsing.
    """
    text = text.strip()

    candidates = [text]
    fence_match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
    if fence_match:
        candidates.append(fence_match.group(1))
    brace_match = re.search(r"\{.*\}", text, re.DOTALL)
    if brace_match:
        candidates.append(brace_match.group(0))

    for candidate in candidates:
        try:
            return _loads_lenient(candidate)
        except json.JSONDecodeError:
            continue

    return None


def judge_single_answer(
    client,
    judge_model_name: str,
    question: str,
    reference_answer: str,
    model_answer: str,
    rpm_delay: float,
) -> dict | None:
    """Send a single judging request to the appropriate LLM provider.

    Dispatches to Gemini or OpenAI based on the model name.

    Returns:
        Parsed score dict or None on failure.
    """
    if _is_openai_model(judge_model_name):
        return _judge_openai(
            client, judge_model_name, question,
            reference_answer, model_answer, rpm_delay,
        )
    return _judge_gemini(
        client, judge_model_name, question,
        reference_answer, model_answer, rpm_delay,
    )


def _validate_scores(scores: dict) -> dict | None:
    """Validate and coerce score types.  Returns dict or None if invalid."""
    if not isinstance(scores, dict):
        return None
    required_keys = {"correctness", "completeness", "clarity",
                     "conciseness", "overall", "justification"}
    if not required_keys.issubset(scores.keys()):
        return None
    try:
        scores["correctness"] = int(scores["correctness"])
        scores["completeness"] = int(scores["completeness"])
        scores["clarity"] = int(scores["clarity"])
        scores["conciseness"] = int(scores["conciseness"])
        scores["overall"] = round(float(scores["overall"]), 2)
        scores["justification"] = str(scores["justification"])
    except (ValueError, TypeError):
        return None
    return scores


def _judge_gemini(
    client: genai.Client,
    judge_model_name: str,
    question: str,
    reference_answer: str,
    model_answer: str,
    rpm_delay: float,
) -> dict | None:
    """Send a single judging request to the Gemini API.

    Returns:
        Parsed score dict or None on failure.
    """
    user_prompt = JUDGE_USER_TEMPLATE.format(
        question=question,
        reference_answer=reference_answer,
        model_answer=model_answer,
    )

    for attempt in range(MAX_RETRIES + 1):
        try:
            time.sleep(rpm_delay)

            response = client.models.generate_content(
                model=judge_model_name,
                contents=user_prompt,
                config=genai_types.GenerateContentConfig(
                    system_instruction=JUDGE_SYSTEM_PROMPT,
                    temperature=0.0,
                ),
            )
            raw_text = response.text

            scores = _parse_judge_json(raw_text)
            if scores is not None:
                validated = _validate_scores(scores)
                if validated is not None:
                    return validated

            if attempt == 0:
                print("    ⚠️  Malformed JSON, retrying...")
                continue
            else:
                print("    ❌ Could not parse judge response after retry")
                return None

        except Exception as e:
            err_str = str(e).lower()
            is_retryable = (
                "429" in err_str
                or "503" in err_str
                or "resource_exhausted" in err_str
                or "resourceexhausted" in err_str
                or "unavailable" in err_str
                or "overloaded" in err_str
                or "quota" in err_str
                or "high demand" in err_str
                or ("internal" in err_str and "500" in err_str)
            )

            if is_retryable and attempt < MAX_RETRIES:
                backoff = min(2 ** (attempt + 1), MAX_BACKOFF_SEC)
                print(f"    ⏳ Retryable error — backing off {backoff}s "
                      f"(attempt {attempt + 1}/{MAX_RETRIES})")
                time.sleep(backoff)
                continue
            else:
                print(f"    ❌ API error (attempt {attempt + 1}/{MAX_RETRIES + 1}): {e}")
                return None

    return None


def _judge_openai(
    client: OpenAI,
    judge_model_name: str,
    question: str,
    reference_answer: str,
    model_answer: str,
    rpm_delay: float,
) -> dict | None:
    """Send a single judging request to the OpenAI Chat Completions API.

    Returns:
        Parsed score dict or None on failure.
    """
    user_prompt = JUDGE_USER_TEMPLATE.format(
        question=question,
        reference_answer=reference_answer,
        model_answer=model_answer,
    )

    for attempt in range(MAX_RETRIES + 1):
        try:
            time.sleep(rpm_delay)

            response = client.chat.completions.create(
                model=judge_model_name,
                messages=[
                    {"role": "system", "content": JUDGE_SYSTEM_PROMPT},
                    {"role": "user",   "content": user_prompt},
                ],
                temperature=0.0,
            )
            raw_text = response.choices[0].message.content or ""

            scores = _parse_judge_json(raw_text)
            if scores is not None:
                validated = _validate_scores(scores)
                if validated is not None:
                    return validated

            if attempt == 0:
                print("    ⚠️  Malformed JSON, retrying...")
                continue
            else:
                print(f"    ❌ Could not parse judge response after retry")
                return None

        except Exception as e:
            err_str = str(e).lower()
            is_retryable = (
                "429" in err_str
                or "503" in err_str
                or "rate_limit" in err_str
                or "rate limit" in err_str
                or "overloaded" in err_str
                or "server_error" in err_str
                or ("internal" in err_str and "500" in err_str)
            )

            if is_retryable and attempt < MAX_RETRIES:
                backoff = min(2 ** (attempt + 1), MAX_BACKOFF_SEC)
                print(f"    ⏳ Retryable error — backing off {backoff}s "
                      f"(attempt {attempt + 1}/{MAX_RETRIES})")
                time.sleep(backoff)
                continue
            else:
                print(f"    ❌ API error (attempt {attempt + 1}/{MAX_RETRIES + 1}): {e}")
                return None

    return None


# ═══════════════════════════════════════════════════════════════════════
# Excel writer — judged results
# ═══════════════════════════════════════════════════════════════════════

def save_judged_to_excel(
    output_path: str,
    model_name: str,
    source_answer_file: str,
    judge_model: str,
    judged_rows: list[dict],
    total_time: float,
    benchmark_file: str = "",
):
    """Save judged results to Excel with Metadata, Judged, and Summary sheets."""
    wb = Workbook()

    scored_rows = [r for r in judged_rows if r.get("overall") is not None]
    failed_rows = [r for r in judged_rows if r.get("overall") is None]

    # ── Metadata sheet ──────────────────────────────────────────────
    ws_meta = wb.active
    ws_meta.title = "Metadata"

    meta_rows = [
        ("Key", "Value"),
        ("model_name", model_name),
        ("source_answer_file", source_answer_file),
        ("judge_model", judge_model),
        ("benchmark_file", benchmark_file),
        ("timestamp", datetime.now().isoformat()),
        ("total_questions", len(judged_rows)),
        ("successfully_judged", len(scored_rows)),
        ("failed_judgments", len(failed_rows)),
        ("total_judging_time_sec", f"{total_time:.1f}"),
        ("avg_judging_time_sec",
         f"{total_time / max(len(scored_rows), 1):.2f}"),
    ]

    for row_idx, (key, value) in enumerate(meta_rows, start=1):
        ws_meta.cell(row=row_idx, column=1, value=key)
        ws_meta.cell(row=row_idx, column=2, value=str(value))
        if row_idx == 1:
            for col in (1, 2):
                ws_meta.cell(row=row_idx, column=col).font = _HEADER_FONT
                ws_meta.cell(row=row_idx, column=col).fill = _HEADER_FILL
                ws_meta.cell(row=row_idx, column=col).alignment = _HEADER_ALIGN
        else:
            ws_meta.cell(row=row_idx, column=1).font = _META_KEY_FONT
            ws_meta.cell(row=row_idx, column=1).border = _THIN_BORDER
            ws_meta.cell(row=row_idx, column=2).border = _THIN_BORDER

    ws_meta.column_dimensions["A"].width = 28
    ws_meta.column_dimensions["B"].width = 80

    # ── Judged sheet ────────────────────────────────────────────────
    ws_judged = wb.create_sheet("Judged")

    headers = [
        "id",
        "category",
        "difficulty",
        "question_type",
        "question",
        "reference_answer",
        "model_answer",
        "correctness",
        "completeness",
        "clarity",
        "conciseness",
        "overall",
        "justification",
        "keywords",
        "overall_formula",
    ]

    for col_idx, header in enumerate(headers, start=1):
        ws_judged.cell(row=1, column=col_idx, value=header)
    _style_header_row(ws_judged, len(headers))

    for row_idx, row_data in enumerate(judged_rows, start=2):
        ws_judged.cell(row=row_idx, column=1, value=row_data.get("id"))
        ws_judged.cell(row=row_idx, column=2, value=row_data.get("category"))
        ws_judged.cell(row=row_idx, column=3, value=row_data.get("difficulty"))
        ws_judged.cell(row=row_idx, column=4, value=row_data.get("question_type", ""))
        ws_judged.cell(row=row_idx, column=5, value=row_data.get("question"))
        ws_judged.cell(row=row_idx, column=6, value=row_data.get("reference_answer"))
        ws_judged.cell(row=row_idx, column=7, value=row_data.get("model_answer"))
        ws_judged.cell(row=row_idx, column=8, value=row_data.get("correctness"))
        ws_judged.cell(row=row_idx, column=9, value=row_data.get("completeness"))
        ws_judged.cell(row=row_idx, column=10, value=row_data.get("clarity"))
        ws_judged.cell(row=row_idx, column=11, value=row_data.get("conciseness"))
        ws_judged.cell(row=row_idx, column=12, value=row_data.get("overall"))
        ws_judged.cell(row=row_idx, column=13, value=row_data.get("justification"))
        ws_judged.cell(row=row_idx, column=14, value=row_data.get("keywords", ""))
        # Deterministic rubric composite alongside the judge-reported overall
        # (docs/JUDGE_OVERALL_AUDIT.md)
        dims = [row_data.get(d) for d in
                ("correctness", "completeness", "clarity", "conciseness")]
        if all(isinstance(v, (int, float)) for v in dims):
            ws_judged.cell(row=row_idx, column=15, value=round(
                0.4 * dims[0] + 0.3 * dims[1] + 0.2 * dims[2] + 0.1 * dims[3], 2))

        for col in range(1, len(headers) + 1):
            cell = ws_judged.cell(row=row_idx, column=col)
            cell.alignment = _WRAP_ALIGN
            cell.border = _THIN_BORDER

        overall_val = row_data.get("overall")
        if overall_val is not None:
            ws_judged.cell(row=row_idx, column=12).fill = _score_fill(
                float(overall_val)
            )

    col_widths = {
        "A": 8,   # id
        "B": 30,  # category
        "C": 12,  # difficulty
        "D": 16,  # question_type
        "E": 55,  # question
        "F": 65,  # reference_answer
        "G": 65,  # model_answer
        "H": 14,  # correctness
        "I": 14,  # completeness
        "J": 10,  # clarity
        "K": 14,  # conciseness
        "L": 10,  # overall
        "M": 60,  # justification
        "N": 35,  # keywords
        "O": 16,  # overall_formula
    }
    for col_letter, width in col_widths.items():
        ws_judged.column_dimensions[col_letter].width = width

    ws_judged.freeze_panes = "A2"
    ws_judged.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(judged_rows) + 1}"
    )

    # ── Summary sheet ───────────────────────────────────────────────
    ws_summary = wb.create_sheet("Summary")
    _build_summary_sheet(ws_summary, scored_rows, model_name)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"💾 Saved: {output_path}")


def _build_summary_sheet(ws, scored_rows: list[dict], model_name: str):
    """Build a summary sheet with aggregate statistics."""
    n = len(scored_rows)
    if n == 0:
        ws.cell(row=1, column=1, value="No successfully judged rows.")
        return

    dims = ["correctness", "completeness", "clarity", "conciseness", "overall"]

    def _avg(rows, key):
        vals = [float(r[key]) for r in rows if r.get(key) is not None]
        return round(sum(vals) / len(vals), 2) if vals else 0.0

    # Section 1: Overall averages
    summary_headers = ["Metric", "Average", "Min", "Max"]
    for col_idx, h in enumerate(summary_headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)
    _style_header_row_at(ws, 1, len(summary_headers))

    for row_off, dim in enumerate(dims, start=2):
        vals = [float(r[dim]) for r in scored_rows if r.get(dim) is not None]
        ws.cell(row=row_off, column=1, value=dim)
        ws.cell(row=row_off, column=2,
                value=round(sum(vals) / len(vals), 2) if vals else "N/A")
        ws.cell(row=row_off, column=3,
                value=round(min(vals), 2) if vals else "N/A")
        ws.cell(row=row_off, column=4,
                value=round(max(vals), 2) if vals else "N/A")
        for col in range(1, 5):
            ws.cell(row=row_off, column=col).border = _THIN_BORDER
        ws.cell(row=row_off, column=1).font = _META_KEY_FONT

    # Section 2: Average by category
    start_row = len(dims) + 3
    ws.cell(row=start_row, column=1, value="Average by Category")
    ws.cell(row=start_row, column=1).font = Font(name="Calibri", bold=True, size=12)

    cat_headers = ["Category", "Count"] + [d.capitalize() for d in dims]
    start_row += 1
    for col_idx, h in enumerate(cat_headers, start=1):
        ws.cell(row=start_row, column=col_idx, value=h)
    _style_header_row_at(ws, start_row, len(cat_headers))

    categories = sorted(set(r.get("category", "Unknown") for r in scored_rows))
    for cat_off, cat in enumerate(categories, start=1):
        r_idx = start_row + cat_off
        cat_rows = [r for r in scored_rows if r.get("category") == cat]
        ws.cell(row=r_idx, column=1, value=cat)
        ws.cell(row=r_idx, column=2, value=len(cat_rows))
        for dim_idx, dim in enumerate(dims, start=3):
            ws.cell(row=r_idx, column=dim_idx, value=_avg(cat_rows, dim))
        for col in range(1, len(cat_headers) + 1):
            ws.cell(row=r_idx, column=col).border = _THIN_BORDER
        ws.cell(row=r_idx, column=1).font = _META_KEY_FONT

    # Section 3: Average by difficulty
    diff_start = start_row + len(categories) + 2
    ws.cell(row=diff_start, column=1, value="Average by Difficulty")
    ws.cell(row=diff_start, column=1).font = Font(
        name="Calibri", bold=True, size=12
    )

    diff_headers = ["Difficulty", "Count"] + [d.capitalize() for d in dims]
    diff_start += 1
    for col_idx, h in enumerate(diff_headers, start=1):
        ws.cell(row=diff_start, column=col_idx, value=h)
    _style_header_row_at(ws, diff_start, len(diff_headers))

    difficulties = ["easy", "medium", "hard"]
    for diff_off, diff in enumerate(difficulties, start=1):
        r_idx = diff_start + diff_off
        diff_rows = [
            r for r in scored_rows
            if str(r.get("difficulty", "")).lower() == diff
        ]
        ws.cell(row=r_idx, column=1, value=diff)
        ws.cell(row=r_idx, column=2, value=len(diff_rows))
        for dim_idx, dim in enumerate(dims, start=3):
            ws.cell(row=r_idx, column=dim_idx, value=_avg(diff_rows, dim))
        for col in range(1, len(diff_headers) + 1):
            ws.cell(row=r_idx, column=col).border = _THIN_BORDER
        ws.cell(row=r_idx, column=1).font = _META_KEY_FONT

    # Column widths
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 12
    for letter in "CDEFG":
        ws.column_dimensions[letter].width = 16


# ═══════════════════════════════════════════════════════════════════════
# Main pipeline — judge one answer file
# ═══════════════════════════════════════════════════════════════════════

def judge_answer_file(
    answer_file: str,
    client,
    judge_model_name: str,
    rpm: int,
    output_dir: str,
):
    """Judge all answers in one Phase 1 Excel file."""
    answer_file_path = Path(answer_file)
    if not answer_file_path.exists():
        print(f"❌ Answer file not found: {answer_file}")
        return

    # Read Phase 1 answers + metadata
    metadata, answers = read_answer_excel(answer_file)

    # Derive model name from metadata (more robust than filename parsing)
    model_name = metadata.get("model_name", "")
    if not model_name:
        fname = answer_file_path.stem
        model_name = fname[len("hpn_answers_"):] if fname.startswith("hpn_answers_") else fname

    # Derive benchmark tag for output naming
    benchmark_file = metadata.get("benchmark_file", "")
    tag = _benchmark_tag(benchmark_file)
    tag_suffix = f"_{tag}" if tag else ""

    print(f"\n{'=' * 70}")
    print(f"⚖️  Judging: {model_name}")
    print(f"   Source: {answer_file}")
    print(f"{'=' * 70}")
    print(f"   📋 Loaded {len(answers)} answers")

    # Output path — include benchmark tag + judge model for distinguishability
    judge_tag = _sanitize_for_filename(judge_model_name)
    output_path = os.path.join(
        output_dir,
        f"hpn_judged_{model_name}{tag_suffix}_by_{judge_tag}.xlsx",
    )

    existing = load_existing_judgments(output_path)
    rpm_delay = 60.0 / rpm

    judged_rows = []
    judged_count = 0
    skipped_count = 0
    failed_count = 0

    total_start = time.perf_counter()

    for idx, ans in enumerate(answers, start=1):
        q_id = str(ans.get("id", ""))

        if q_id in existing:
            judged_rows.append(existing[q_id])
            skipped_count += 1
            continue

        question = str(ans.get("question", ""))
        ref_answer = str(ans.get("reference_answer", ""))
        model_answer = str(ans.get("model_answer", ""))

        print(f"  [{idx:3d}/{len(answers)}] {q_id} ...", end=" ", flush=True)

        scores = judge_single_answer(
            client=client,
            judge_model_name=judge_model_name,
            question=question,
            reference_answer=ref_answer,
            model_answer=model_answer,
            rpm_delay=rpm_delay,
        )

        row = {
            "id": ans.get("id"),
            "category": ans.get("category"),
            "difficulty": ans.get("difficulty"),
            "question_type": ans.get("question_type", ""),
            "question": question,
            "reference_answer": ref_answer,
            "model_answer": model_answer,
            "keywords": ans.get("keywords", ""),
        }

        if scores:
            row.update(scores)
            judged_count += 1
            print(f"✅ overall={scores['overall']}")
        else:
            row.update({
                "correctness": None,
                "completeness": None,
                "clarity": None,
                "conciseness": None,
                "overall": None,
                "justification": "JUDGE FAILED — no valid score returned",
            })
            failed_count += 1
            print("❌ FAILED")

        judged_rows.append(row)

    total_time = time.perf_counter() - total_start

    save_judged_to_excel(
        output_path=output_path,
        model_name=model_name,
        source_answer_file=answer_file,
        judge_model=judge_model_name,
        judged_rows=judged_rows,
        total_time=total_time,
        benchmark_file=benchmark_file,
    )

    scored_rows = [r for r in judged_rows if r.get("overall") is not None]
    avg_overall = (
        sum(float(r["overall"]) for r in scored_rows) / len(scored_rows)
        if scored_rows else 0.0
    )

    print(f"\n{'─' * 70}")
    print(f"✅ {model_name} — judging complete")
    print(f"   Newly judged : {judged_count}")
    print(f"   Resumed      : {skipped_count}")
    print(f"   Failed       : {failed_count}")
    print(f"   Total scored : {len(scored_rows)} / {len(judged_rows)}")
    print(f"   Avg overall  : {avg_overall:.2f}")
    print(f"   Time (new)   : {total_time:.1f}s")
    print(f"   Output       : {output_path}")
    print(f"{'─' * 70}\n")


# ═══════════════════════════════════════════════════════════════════════
# CLI
# ═══════════════════════════════════════════════════════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="RAG Pipeline Benchmark — LLM-as-Judge (Phase 2)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Examples:\n"
            "  # Judge with Gemini (default)\n"
            "  python evaluation/judge_responses.py \\\n"
            "      --answer_files outputs/answers/hpn_answers_RAG-gpt-4o.xlsx\n"
            "\n"
            "  # Judge with OpenAI GPT-4o\n"
            "  python evaluation/judge_responses.py \\\n"
            "      --answer_files outputs/answers/hpn_answers_RAG-gpt-4o.xlsx \\\n"
            "      --judge_model gpt-4o \\\n"
            "      --openai_api_key_file openai_api_key.txt\n"
            "\n"
            "  # Custom rate limit and Gemini judge model\n"
            "  python evaluation/judge_responses.py \\\n"
            "      --answer_files outputs/answers/hpn_answers_RAG-gpt-4o.xlsx \\\n"
            "      --judge_model gemini-2.5-pro \\\n"
            "      --requests_per_minute 30\n"
        ),
    )

    parser.add_argument(
        "--answer_files",
        type=str,
        nargs="+",
        required=True,
        help="Path(s) to Phase 1 answer Excel file(s)",
    )
    parser.add_argument(
        "--config",
        type=str,
        default=DEFAULT_CONFIG,
        help=f"Path to config.yaml (for API key and judge settings) "
             f"(default: {DEFAULT_CONFIG})",
    )
    parser.add_argument(
        "--gemini_api_key_file",
        type=str,
        default=None,
        help="Path to a text file containing the Gemini API key "
             "(overrides config.yaml and env var)",
    )
    parser.add_argument(
        "--openai_api_key_file",
        type=str,
        default=None,
        help="Path to a text file containing the OpenAI API key "
             "(overrides config.yaml and env var)",
    )
    parser.add_argument(
        "--judge_model",
        type=str,
        default=None,
        help=f"Model to use as judge "
             f"(default: from config.yaml judge.model or {DEFAULT_JUDGE_MODEL}). "
             "Gemini models: gemini-2.5-flash, gemini-2.5-pro, etc. "
             "OpenAI models: gpt-4o, gpt-4o-mini, gpt-4.1, etc.",
    )
    parser.add_argument(
        "--requests_per_minute",
        type=int,
        default=None,
        help=f"Max API requests per minute "
             f"(default: from config.yaml judge.rate_limit_rpm or {DEFAULT_RPM})",
    )
    parser.add_argument(
        "--output_dir",
        type=str,
        default=DEFAULT_OUTPUT_DIR,
        help=f"Directory for output Excel files (default: {DEFAULT_OUTPUT_DIR})",
    )

    args = parser.parse_args()

    # ── Load config ─────────────────────────────────────────────────
    cfg = {}
    config_path = Path(args.config)
    if config_path.exists():
        with open(config_path, encoding="utf-8") as f:
            cfg = yaml.safe_load(f) or {}

    judge_cfg = cfg.get("judge", {})
    judge_model = (
        args.judge_model
        or judge_cfg.get("model", DEFAULT_JUDGE_MODEL)
    )
    rpm = (
        args.requests_per_minute
        or judge_cfg.get("rate_limit_rpm", DEFAULT_RPM)
    )

    # ── Banner ──────────────────────────────────────────────────────
    print("\n" + "=" * 70)
    print("  ⚖️   RAG Pipeline Benchmark — LLM-as-Judge (Phase 2)")
    print("=" * 70)

    # ── Load API key & create client (provider-dependent) ─────────
    use_openai = _is_openai_model(judge_model)
    if use_openai:
        api_key = load_openai_api_key(args.openai_api_key_file, cfg)
        client = OpenAI(api_key=api_key)
        provider_label = "OpenAI"
    else:
        api_key = load_api_key(args.gemini_api_key_file, cfg)
        client = genai.Client(api_key=api_key)
        provider_label = "Gemini"

    print(f"\n⚙️  Configuration:")
    print(f"   Provider         : {provider_label}")
    print(f"   Judge model      : {judge_model}")
    print(f"   Requests/min     : {rpm}")
    print(f"   Output dir       : {args.output_dir}")
    print(f"   Answer files     : {len(args.answer_files)}")
    for i, af in enumerate(args.answer_files, 1):
        print(f"     {i}. {af}")

    # ── Judge each answer file ──────────────────────────────────────
    overall_start = time.perf_counter()

    for answer_file in args.answer_files:
        judge_answer_file(
            answer_file=answer_file,
            client=client,
            judge_model_name=judge_model,
            rpm=rpm,
            output_dir=args.output_dir,
        )

    overall_time = time.perf_counter() - overall_start

    print("=" * 70)
    print(f"🏁 All done! {len(args.answer_files)} file(s) judged "
          f"in {overall_time:.1f}s")
    print(f"   Output files in: {args.output_dir}/")
    print(f"   (Exact filenames depend on benchmark tag in each answer file's metadata.)")
    print("=" * 70 + "\n")


if __name__ == "__main__":
    main()
