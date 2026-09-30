"""Re-run retrieval alone, logging chunk provenance the answer workbooks lost.

Phase 12 R5. The committed RAG workbooks record each retrieved chunk as a
display label -- `paper_title[:40] + "… §" + section` (evaluate_rag.py:499-503)
-- which is enough to bracket gold-evidence recall but not to measure it:
distinct papers collide in the first 40 characters, and a label carries no rank
beyond the final four. This script measures it directly.

Retrieval only. No generation, no judging, no API key, no training. The indexes
are read as they stand and nothing is rewritten, so this does not disturb the
committed evidence chain -- it adds a provenance log beside it.

Two properties make one pass sufficient, both verified rather than assumed:

  * Retrieval is deterministic and does not depend on the generator: every one
    of the committed RAG runs returns byte-identical `retrieved_chunks` for
    every question. So a single pass reproduces what all of them retrieved.
  * The cross-encoder already scores all `rrf_top_n` candidates and only then
    truncates to `reranker_top_n` (retriever.py:200-206). Raising the cutoff to
    the fusion width therefore costs nothing and yields the full ranking, whose
    top-4 prefix is exactly what the original runs consumed. `--verify` checks
    that prefix against a committed workbook rather than trusting the argument.

    cd NetBench-RAG && venv/bin/python evaluation/retrieval_only.py \\
        --verify outputs/answers/hpn_answers_RAG-Qwen3.5-9B_v5.0.xlsx
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import time

import yaml

sys.path.insert(0, os.getcwd())
from src.retriever import retriever_from_config          # noqa: E402


def chunk_label(c) -> str:
    """evaluate_rag.py's display label, rebuilt for the verification check."""
    t = c.paper_title[:40] + "…" if len(c.paper_title) > 40 else c.paper_title
    sec = (f"§{c.section_number} {c.section_heading}".strip()
           if c.section_number else f"§{c.section_heading}")
    return f"{t} {sec}"


def load_questions(path: str) -> list[dict]:
    with open(path, encoding="utf-8") as fh:
        d = json.load(fh)
    for key in ("questions", "items", "data"):
        if isinstance(d, dict) and key in d:
            return d[key]
    return d if isinstance(d, list) else []


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--config", default="config.yaml")
    ap.add_argument("--out", default="outputs/retrieval_provenance.csv")
    ap.add_argument("--verify", default=None,
                    help="committed answer workbook to check the top-4 against")
    args = ap.parse_args()

    cfg = yaml.safe_load(open(args.config, encoding="utf-8"))
    depth = cfg["retrieval"]["rrf_top_n"]
    original_top_n = cfg["retrieval"]["reranker_top_n"]
    # Rank the whole fusion candidate set instead of only the served prefix.
    # This is a reporting change, not a retrieval change: the same candidates
    # get the same scores either way.
    cfg["retrieval"]["reranker_top_n"] = depth

    questions = load_questions(cfg["paths"]["benchmark"])
    print(f"{len(questions)} questions; ranking all {depth} fusion candidates "
          f"(runs served the top {original_top_n})")

    t0 = time.perf_counter()
    retriever = retriever_from_config(cfg)
    print(f"retriever loaded in {time.perf_counter() - t0:.1f}s")

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    served: dict[str, list[str]] = {}
    t0 = time.perf_counter()
    with open(args.out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["question_id", "rank", "source_file", "paper_id",
                    "section_number", "section_heading", "chunk_index",
                    "reranker_score", "label"])
        for i, q in enumerate(questions, 1):
            qid = q.get("id") or q.get("question_id")
            res = retriever.retrieve(q["question"])
            for rank, (c, s) in enumerate(zip(res.chunks, res.reranker_scores), 1):
                w.writerow([qid, rank, c.source_file,
                            "paper_" + str(c.source_file).rsplit("_", 1)[-1],
                            c.section_number, c.section_heading, c.chunk_index,
                            f"{s:.6f}", chunk_label(c)])
            served[qid] = [chunk_label(c) for c in res.chunks[:original_top_n]]
            if i % 40 == 0:
                print(f"  {i}/{len(questions)}  "
                      f"{(time.perf_counter() - t0) / i:.2f}s/question")
    dt = time.perf_counter() - t0
    print(f"retrieved {len(questions)} questions in {dt:.0f}s "
          f"({dt / len(questions):.2f}s each) -> {args.out}")

    if args.verify:
        # openpyxl, not pandas: this runs in the NetBench-RAG venv, which has
        # the former (evaluate_rag.py writes these workbooks with it) and not
        # the latter.
        from openpyxl import load_workbook
        ws = load_workbook(args.verify, read_only=True)["Answers"]
        head = [c.value for c in next(ws.rows)]
        i_id, i_ch = head.index("id"), head.index("retrieved_chunks")
        same = diff = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            got, logged = served.get(row[i_id]), row[i_ch]
            if got is None or not isinstance(logged, str):
                continue
            same += got == logged.split(" | ")
            diff += got != logged.split(" | ")
        print(f"verify against {os.path.basename(args.verify)}: "
              f"{same} questions reproduce the committed top-{original_top_n} "
              f"exactly, {diff} differ")


if __name__ == "__main__":
    main()
